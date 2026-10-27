import csv
from io import BytesIO, StringIO
from zipfile import BadZipFile

from django.conf import settings
from django.core.exceptions import ValidationError

from .forms import validate_unique_username
from .models import Course


IMPORT_HEADERS = ('first_name', 'last_name', 'username', 'gender', 'course')


class LearnerImportError(ValueError):
    pass


def parse_learner_upload(uploaded_file):
    extension = uploaded_file.name.rsplit('.', 1)[-1].lower()
    contents = uploaded_file.read()
    if len(contents) > settings.SIMS_LEARNER_IMPORT_MAX_BYTES:
        maximum_megabytes = (
            settings.SIMS_LEARNER_IMPORT_MAX_BYTES / (1024 * 1024))
        raise LearnerImportError(
            f'The upload exceeds the {maximum_megabytes:g} MB file-size limit.')

    if extension == 'csv':
        try:
            text = contents.decode('utf-8-sig')
        except UnicodeDecodeError as error:
            raise LearnerImportError(
                'The CSV file must use UTF-8 encoding.') from error
        try:
            raw_rows = list(csv.reader(StringIO(text), strict=True))
        except csv.Error as error:
            raise LearnerImportError(
                'The CSV file could not be parsed.') from error
    elif extension == 'xlsx':
        raw_rows = _read_xlsx(contents)
    else:
        raise LearnerImportError('Choose a .csv or .xlsx file.')

    if not raw_rows:
        raise LearnerImportError('The file is empty.')
    headers = [
        str(value or '').strip().lower()
        for value in raw_rows[0]
    ]
    if len(headers) != len(set(headers)):
        raise LearnerImportError('The file contains duplicate column names.')
    missing = [header for header in IMPORT_HEADERS if header not in headers]
    unexpected = [header for header in headers if header not in IMPORT_HEADERS]
    if missing or unexpected:
        details = []
        if missing:
            details.append('Missing columns: ' + ', '.join(missing))
        if unexpected:
            details.append('Unexpected columns: ' + ', '.join(unexpected))
        raise LearnerImportError(
            '; '.join(details) + '. Use the provided import template.')

    column_indexes = [headers.index(header) for header in IMPORT_HEADERS]
    data_rows = [
        (line_number, row)
        for line_number, row in enumerate(raw_rows[1:], start=2)
        if any(str(value or '').strip() for value in row)
    ]
    if not data_rows:
        raise LearnerImportError('The file contains no learner rows.')
    if len(data_rows) > settings.SIMS_LEARNER_IMPORT_MAX_ROWS:
        raise LearnerImportError(
            'The file exceeds the '
            f'{settings.SIMS_LEARNER_IMPORT_MAX_ROWS}-row limit.')

    candidate_rows = []
    for line_number, row in data_rows:
        if len(row) > len(headers) and any(
            str(value or '').strip() for value in row[len(headers):]
        ):
            candidate_rows.append({
                'line': line_number,
                'data': {},
                'errors': ['The row has values beyond the listed columns.'],
            })
            continue
        values = {
            header: (
                str(row[index] or '').strip()
                if index < len(row) else ''
            )
            for header, index in zip(IMPORT_HEADERS, column_indexes)
        }
        candidate_rows.append({
            'line': line_number,
            'data': values,
            'errors': [],
        })
    return validate_learner_rows(candidate_rows)


def validate_learner_rows(candidate_rows):
    seen_usernames = set()
    validated_rows = []
    for candidate in candidate_rows:
        data = dict(candidate.get('data', {}))
        errors = list(candidate.get('errors', []))
        for field in ('first_name', 'last_name'):
            if not data.get(field):
                errors.append(f'{field.replace("_", " ").title()} is required.')
            elif len(data[field]) > 150:
                errors.append(
                    f'{field.replace("_", " ").title()} must be 150 '
                    'characters or fewer.')

        username = data.get('username', '')
        if username:
            try:
                username = validate_unique_username(username)
            except ValidationError as error:
                errors.extend(error.messages)
            key = username.casefold()
            if key in seen_usernames:
                errors.append('Username is duplicated in this file.')
            seen_usernames.add(key)
        else:
            errors.append('Username is required.')

        gender = data.get('gender', '').strip().upper()
        if gender not in ('M', 'F'):
            errors.append('Gender must be M or F.')

        course = None
        course_name = data.get('course', '')
        if course_name:
            course = Course.objects.filter(
                name__iexact=course_name).order_by('pk').first()
            if course is None:
                errors.append(f'Course / grade "{course_name}" does not exist.')

        data['username'] = username
        data['gender'] = gender
        data['course_id'] = course.pk if course else None
        data['course_name'] = course.name if course else ''
        validated_rows.append({
            'line': candidate['line'],
            'data': data,
            'errors': errors,
        })
    return validated_rows


def _read_xlsx(contents):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:
        raise LearnerImportError(
            'XLSX import is unavailable because its spreadsheet dependency '
            'is not installed.') from error

    try:
        workbook = load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=False,
        )
    except (
        OSError, ValueError, KeyError, BadZipFile, InvalidFileException,
    ) as error:
        raise LearnerImportError('The XLSX file could not be read.') from error

    worksheet = workbook.active
    raw_rows = []
    try:
        for row_number, cells in enumerate(
            worksheet.iter_rows(values_only=False), start=1
        ):
            values = []
            for cell in cells:
                if cell.data_type == 'f':
                    raise LearnerImportError(
                        f'Formula cells are not allowed (row {row_number}).')
                values.append(cell.value)
            raw_rows.append(values)
            if row_number > settings.SIMS_LEARNER_IMPORT_MAX_ROWS + 1:
                raise LearnerImportError(
                    'The file exceeds the '
                    f'{settings.SIMS_LEARNER_IMPORT_MAX_ROWS}-row limit.')
    finally:
        workbook.close()
    return raw_rows
