import csv
import tempfile
from io import BytesIO, StringIO
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.core.cache import cache
from django.contrib.auth.hashers import check_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.http import HttpResponse
from django.shortcuts import render
from django.test import RequestFactory, TestCase
from django.test.utils import override_settings
from django.templatetags.static import static
from django.urls import Resolver404, resolve, reverse
from django.utils import timezone
from PIL import Image

from .models import (
    Admin,
    ActivityLog,
    Attendance,
    AttendanceReport,
    Course,
    CustomUser,
    GradeBand,
    LoginActivity,
    AcademicPeriod,
    Student,
    StudentResult,
    Subject,
    Teacher,
)
from .credentials import generate_pin, issue_pin, validate_pin
from .views import render_pin_confirmation
from .activity import reset_current_request, set_current_request


class CrestBrandingTests(TestCase):
    def _assert_page_branding(self, response):
        self.assertContains(
            response,
            f'src="{static("main_app/img/crest-logo.png")}"',
        )
        self.assertContains(
            response,
            f'href="{static("main_app/img/favicon.ico")}"',
        )
        self.assertContains(
            response,
            f'href="{static("main_app/img/favicon-32x32.png")}"',
        )
        self.assertContains(
            response,
            f'href="{static("main_app/img/favicon-192x192.png")}"',
        )
        self.assertContains(
            response,
            f'href="{static("main_app/img/apple-touch-icon.png")}"',
        )
        self.assertContains(response, 'id="page-loader"')
        self.assertContains(
            response,
            static('main_app/css/page-loader.css'),
        )
        self.assertContains(
            response,
            static('main_app/js/page-loader.js'),
        )

    def test_landing_sign_in_learner_registration_and_pin_pages_show_branding(self):
        self._assert_page_branding(self.client.get(reverse('login_page')))
        for role in ('admin', 'teacher', 'learner'):
            with self.subTest(role=role):
                self._assert_page_branding(
                    self.client.get(reverse('role_login', args=[role])))
        self._assert_page_branding(
            self.client.get(reverse('register_learner')))

        pin_response = render_pin_confirmation(
            RequestFactory().get('/'),
            'learner',
            '12345',
        )
        self._assert_page_branding(pin_response)
        account = CustomUser.objects.create_user(
            username='branding-account',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        reset_response = render(
            RequestFactory().get('/'),
            'main_app/credential_reset_confirmation.html',
            {
                'account': account,
                'secret': 'new-secret',
                'secret_label': 'New password',
                'return_url': '/',
            },
        )
        self._assert_page_branding(reset_response)

    def test_all_three_panels_show_the_crest_and_role_name(self):
        panels = (
            ('1', 'admin_home', 'Admin Panel', 'admin-branding'),
            ('3', 'teacher_home', 'Teacher Panel', 'teacher-branding'),
            ('2', 'student_home', 'Learner Panel', 'learner-branding'),
        )
        for role, page, panel_name, username in panels:
            with self.subTest(role=role):
                user = CustomUser.objects.create_user(
                    username=username,
                    password='valid-password',
                    user_type=role,
                    gender='F',
                    address='',
                )
                self.client.force_login(user)
                response = self.client.get(reverse(page))
                self.assertEqual(response.status_code, 200)
                self.assertContains(
                    response, 'alt="C. Heuva school crest"')
                self.assertContains(response, 'id="page-loader"')
                self.assertContains(response, panel_name)
                self.client.logout()

    @override_settings(DEBUG=True)
    def test_static_branding_assets_are_served_and_favicon_is_public(self):
        image_names = (
            'crest-logo.png',
            'favicon.ico',
            'favicon-32x32.png',
            'favicon-192x192.png',
            'apple-touch-icon.png',
        )
        for image_name in image_names:
            with self.subTest(image=image_name):
                url = static(f'main_app/img/{image_name}')
                response = self.client.get(url)
                self.assertEqual(
                    response.status_code, 200,
                    f'{url} resolved as {response.wsgi_request.resolver_match}')
                expected = Path(
                    'main_app/static/main_app/img', image_name).read_bytes()
                self.assertEqual(b''.join(response.streaming_content), expected)

        for asset in (
                'main_app/css/page-loader.css',
                'main_app/js/page-loader.js',
        ):
                with self.subTest(asset=asset):
                    response = self.client.get(static(asset))
                    self.assertEqual(response.status_code, 200)
                    self.assertTrue(b''.join(response.streaming_content))

        response = self.client.get('/favicon.ico', follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.redirect_chain[0][0],
            static('main_app/img/favicon.ico'),
        )
        self.assertEqual(
            b''.join(response.streaming_content),
            Path('main_app/static/main_app/img/favicon.ico').read_bytes(),
        )

    def test_reports_include_crest_and_still_render_without_it(self):
        admin = CustomUser.objects.create_user(
            username='crest-report-admin',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        from .hod_views import _create_report_pdf

        response = _create_report_pdf(
            admin,
            'Crest Test Report',
            ['Column'],
            [['Value']],
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(b'/Subtype /Image', response.content)

        with tempfile.TemporaryDirectory() as static_root:
            with patch(
                    'main_app.hod_views.finders.find',
                    return_value=None), patch(
                    'main_app.hod_views.settings.STATIC_ROOT',
                    static_root):
                with self.assertLogs('main_app.hod_views', level='WARNING'):
                    response = _create_report_pdf(
                        admin,
                        'Report without crest',
                        ['Column'],
                        [['Value']],
                    )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertNotIn(b'/Subtype /Image', response.content)


class InterfaceSelectionTests(TestCase):
    @staticmethod
    def _profile_image_upload():
        image_data = BytesIO()
        Image.new('RGB', (1, 1), color='white').save(image_data, format='PNG')
        return SimpleUploadedFile(
            'profile.png', image_data.getvalue(), content_type='image/png')

    def test_landing_page_offers_all_three_interfaces(self):
        response = self.client.get(reverse('login_page'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Administrator')
        self.assertContains(response, 'Teacher')
        self.assertContains(response, 'Learner')
        self.assertContains(response, reverse('role_login', args=['admin']))
        self.assertContains(response, reverse('role_login', args=['teacher']))
        self.assertContains(response, reverse('role_login', args=['learner']))

    def test_each_role_has_a_login_page(self):
        for role in ('admin', 'teacher', 'learner'):
            with self.subTest(role=role):
                response = self.client.get(reverse('role_login', args=[role]))

                self.assertEqual(response.status_code, 200)
                self.assertContains(response, f'Sign in as {role.title()}')
                if role == 'learner':
                    self.assertContains(response, 'name=\'password\'')
                    self.assertNotContains(response, 'name="pin"')
                    self.assertContains(
                        response, "Don't have an account? Register")
                    self.assertContains(response, 'Choose a different interface')
                else:
                    self.assertContains(response, 'name="pin"')
                    self.assertNotContains(response, 'name=\'password\'')
                    self.assertContains(response, 'Choose a different interface')
                    self.assertContains(response, 'data-toggle-pin')
                    self.assertContains(response, 'aria-label="Show PIN"')

    def test_learner_can_sign_in_through_learner_interface(self):
        learner = CustomUser.objects.create_user(
            email='learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )

        response = self.client.post(reverse('user_login'), {
            'username': 'learner@example.com',
            'password': 'valid-password',
            'role': 'learner',
        }, REMOTE_ADDR='203.0.113.12')

        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )
        activity = LoginActivity.objects.get(student=learner.student)
        self.assertEqual(activity.ip_address, '203.0.113.12')
        self.assertFalse(learner.pin_hash)
        self.assertTrue(learner.check_password('valid-password'))

    def test_administrator_login_does_not_create_login_activity(self):
        admin = CustomUser.objects.create_user(
            username='admin-login',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        admin_pin = issue_pin(admin)

        response = self.client.post(reverse('user_login'), {
            'username': admin.username,
            'password': 'valid-password',
            'role': 'admin',
        })
        self.assertRedirects(
            response,
            reverse('role_login', args=['admin']),
            fetch_redirect_response=False,
        )
        response = self.client.post(reverse('user_login'), {
            'username': admin.username,
            'pin': admin_pin,
            'role': 'admin',
        })

        self.assertRedirects(
            response,
            reverse('admin_home'),
            fetch_redirect_response=False,
        )
        self.assertEqual(LoginActivity.objects.count(), 0)

    def test_failed_learner_login_does_not_create_login_activity(self):
        learner = CustomUser.objects.create_user(
            email='failed-login@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )

        response = self.client.post(reverse('user_login'), {
            'username': 'failed-login@example.com',
            'password': 'wrong-password',
            'role': 'learner',
        })

        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
            fetch_redirect_response=False,
        )
        self.assertEqual(LoginActivity.objects.count(), 0)
        failure = ActivityLog.objects.get(
            action='failed_sign_in',
            object_type='authentication',
            target_user=learner,
        )
        self.assertNotIn('wrong-password', failure.description)
        self.assertEqual(failure.changes, {})

    def test_learner_profile_is_created_and_dashboard_loads(self):
        learner = CustomUser.objects.create_user(
            email='dashboard@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )

        self.client.force_login(learner)
        response = self.client.get(reverse('student_home'))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(learner.student.pk)

    def test_learner_dashboard_links_to_results_and_available_report_card(self):
        learner = CustomUser.objects.create_user(
            email='dashboard-report-card@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        learner.student.academic_period = academic_period
        learner.student.save()
        subject = Subject.objects.create(name='Mathematics')
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=20, exam=70)
        self.client.force_login(learner)

        response = self.client.get(reverse('student_home'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'My Results')
        self.assertContains(response, reverse('student_view_result'))
        self.assertContains(response, 'Download report card')
        self.assertContains(response, reverse('student_report_card'))

    def test_learner_dashboard_shows_report_card_availability_when_no_results(self):
        learner = CustomUser.objects.create_user(
            email='dashboard-no-report@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('student_home'))

        self.assertContains(response, 'My Report Card')
        self.assertContains(
            response,
            'Available after an academic period / term and results are assigned.',
        )
        self.assertNotContains(response, reverse('student_report_card'))

    def test_learner_dashboard_loads_without_attendance_or_course_features(self):
        learner = CustomUser.objects.create_user(
            email='unassigned@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('student_home'))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'View Attendance')
        self.assertNotContains(response, 'Total Attendance')

    def test_learner_sees_only_subjects_assigned_to_their_course(self):
        learner = CustomUser.objects.create_user(
            email='subjects-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        learner_course = Course.objects.create(name='Grade 7')
        other_course = Course.objects.create(name='Grade 8')
        learner.student.course = learner_course
        learner.student.save()
        Subject.objects.create(name='Mathematics', course=learner_course)
        Subject.objects.create(name='History', course=other_course)
        self.client.force_login(learner)

        response = self.client.get(reverse('student_view_subjects'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Mathematics')
        self.assertNotContains(response, 'History')
        self.assertNotContains(response, 'Add Subject')

    def test_subject_added_by_admin_appears_on_learner_subject_page(self):
        admin = CustomUser.objects.create_user(
            email='subject-author@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='subject-viewer@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(admin)
        add_response = self.client.post(reverse('add_subject'), {
            'name': 'Science',
        })
        self.assertRedirects(
            add_response, reverse('manage_subject'), fetch_redirect_response=False)

        self.client.force_login(learner)
        learner_response = self.client.get(reverse('student_view_subjects'))

        self.assertEqual(learner_response.status_code, 200)
        self.assertContains(learner_response, 'Science')

    def test_learner_subject_page_shows_message_when_no_subjects_are_assigned(self):
        learner = CustomUser.objects.create_user(
            email='no-subjects-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('student_view_subjects'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'No subjects have been added by the administrator yet.',
        )

    def test_admin_dashboard_excludes_teacher_managed_tools(self):
        admin = CustomUser.objects.create_user(
            email='dashboard-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.get(reverse('admin_home'))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Grade Bands')
        self.assertNotContains(response, 'Learner Report Cards')
        self.assertNotContains(response, reverse('generate_report_card'))
        for label in (
            'Add Course',
            'Manage Course',
            'Add Academic Period',
            'Manage Academic Period',
            'View Attendance',
        ):
            with self.subTest(label=label):
                self.assertNotContains(response, label)

    def test_manage_students_has_only_add_student_modal(self):
        admin = CustomUser.objects.create_user(
            email='add-student-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        CustomUser.objects.create_user(
            username='listed-learner',
            email='legacy-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.get(reverse('manage_student'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-target="#addStudentModal"')
        self.assertContains(response, 'action="' + reverse('add_student') + '"')
        self.assertNotContains(response, 'addSubjectModal')
        self.assertContains(response, 'id_course')
        self.assertContains(response, 'id_academic_period')
        self.assertContains(response, 'Add Learner')
        self.assertNotContains(response, 'Add Student')
        self.assertNotContains(response, 'Add Subject</p>')
        self.assertContains(response, 'listed-learner')
        self.assertNotContains(response, 'legacy-learner@example.com')

    def test_address_and_profile_picture_are_optional_in_staff_forms(self):
        from .forms import StudentForm, TeacherAccountForm

        student_form = StudentForm()
        teacher_form = TeacherAccountForm()

        self.assertFalse(student_form.fields['address'].required)
        self.assertFalse(student_form.fields['profile_pic'].required)
        self.assertFalse(teacher_form.fields['profile_pic'].required)

    def test_manage_students_restores_missing_learner_profile(self):
        admin = CustomUser.objects.create_user(
            email='missing-profile-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='missing-profile-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        learner.student.delete()
        self.client.force_login(admin)

        response = self.client.get(reverse('manage_student'))

        self.assertEqual(response.status_code, 200)
        learner_profile = Student.objects.get(admin=learner)
        self.assertContains(
            response, reverse('edit_student', args=[learner_profile.id]))
        self.assertTrue(Student.objects.filter(admin=learner).exists())

    def test_add_pages_redirect_to_manage_students(self):
        admin = CustomUser.objects.create_user(
            email='add-subject-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        student_response = self.client.get(reverse('add_student'))
        subject_response = self.client.get(reverse('add_subject'))

        self.assertRedirects(
            student_response, reverse('manage_student'),
            fetch_redirect_response=False)
        self.assertRedirects(
            subject_response, reverse('manage_subject'),
            fetch_redirect_response=False)

    def test_manage_subjects_has_add_subject_modal(self):
        admin = CustomUser.objects.create_user(
            email='manage-subject-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.get(reverse('manage_subject'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-target="#addSubjectModal"')
        self.assertContains(response, 'action="' + reverse('add_subject') + '"')
        self.assertContains(response, 'id_course')

    def test_invalid_subject_submission_reopens_subject_modal(self):
        admin = CustomUser.objects.create_user(
            email='invalid-subject-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.post(reverse('add_subject'), {'name': ''})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "$('#addSubjectModal').modal('show')")
        self.assertContains(response, 'This field is required.')

    def test_navigation_shows_each_role_only_its_functions(self):
        admin = CustomUser.objects.create_user(
            email='report-sidebar-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='report-sidebar-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        teacher = CustomUser.objects.create_user(
            email='report-sidebar-teacher@example.com',
            password='valid-password',
            user_type='3',
            gender='F',
            address='',
        )

        self.client.force_login(admin)
        admin_response = self.client.get(reverse('admin_home'))
        self.assertContains(admin_response, 'School Summary / Class List')
        self.assertNotContains(admin_response, 'Manage Results')
        self.assertNotContains(admin_response, reverse('manage_student_results'))
        self.assertNotContains(admin_response, 'Grade Bands')
        self.assertNotContains(admin_response, reverse('generate_report_card'))
        self.assertContains(
            admin_response, reverse('generate_report'), html=False)
        admin_html = admin_response.content.decode()
        self.assertLess(
            admin_html.index('School Summary / Class List'),
            admin_html.index('Logout'),
        )
        self.assertIn('fas fa-home', admin_html)
        self.assertIn('fas fa-user-edit', admin_html)
        self.assertIn('fas fa-calendar-minus', admin_html)
        self.assertIn('fas fa-comments', admin_html)

        self.client.force_login(learner)
        learner_response = self.client.get(reverse('student_home'))
        self.assertNotContains(learner_response, 'Generate Report')
        learner_html = learner_response.content.decode()
        self.assertIn('fas fa-home', learner_html)
        self.assertIn('fas fa-user-edit', learner_html)
        self.assertIn('fas fa-calendar-minus', learner_html)
        self.assertIn('fas fa-comments', learner_html)

        self.client.force_login(teacher)
        teacher_response = self.client.get(reverse('teacher_home'))
        self.assertContains(teacher_response, 'Manage Results')
        self.assertNotContains(teacher_response, 'Grade Bands')
        self.assertNotContains(teacher_response, reverse('manage_teachers'))
        self.assertNotContains(teacher_response, reverse('learner_login_activity'))

    def test_admin_search_bars_render_on_requested_pages(self):
        admin = CustomUser.objects.create_user(
            email='search-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)
        pages = (
            ('admin_notify_student', 'notify-student-list'),
            ('manage_student', 'manage-student-list'),
            ('manage_subject', 'manage-subject-list'),
            ('student_feedback_message', 'admin-feedback-list'),
            ('view_student_leave', 'admin-leave-list'),
        )
        for page_name, target_id in pages:
            with self.subTest(page=page_name):
                response = self.client.get(reverse(page_name))
                self.assertEqual(response.status_code, 200)
                self.assertContains(
                    response, f'data-search-target="#{target_id}"')

    def test_learner_search_bars_render_on_requested_pages(self):
        learner = CustomUser.objects.create_user(
            email='search-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)
        pages = (
            ('student_view_subjects', 'learner-subject-list'),
            ('student_view_result', 'learner-results-list'),
            ('student_view_notification', 'learner-notifications-list'),
            ('student_apply_leave', 'learner-leave-list'),
            ('student_feedback', 'learner-feedback-list'),
        )
        for page_name, target_id in pages:
            with self.subTest(page=page_name):
                response = self.client.get(reverse(page_name))
                self.assertEqual(response.status_code, 200)
                self.assertContains(
                    response, f'data-search-target="#{target_id}"')

    def test_learner_terminology_is_used_in_role_and_admin_labels(self):
        admin = CustomUser.objects.create_user(
            email='learner-terminology-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)
        response = self.client.get(reverse('admin_home'))

        self.assertContains(response, 'Manage Learners')
        self.assertContains(response, 'Total Learners')
        self.assertEqual(str(Student._meta.verbose_name), 'learner')
        self.assertEqual(
            str(StudentResult._meta.get_field('student').verbose_name),
            'learner')
        self.assertEqual(CustomUser.USER_TYPE[1][1], 'Learner')

    def test_removed_management_and_attendance_routes_are_not_registered(self):
        removed_paths = (
            '/course/add',
            '/course/manage/',
            '/add_session/',
            '/academic_period/manage/',
            '/attendance/view/',
            '/attendance/fetch/',
            '/get_attendance',
            '/student/view/attendance/',
        )
        for path in removed_paths:
            with self.subTest(path=path):
                with self.assertRaises(Resolver404):
                    resolve(path)

    def test_student_and_subject_forms_support_report_class_filters(self):
        from .forms import StudentForm

        self.assertIn('course', StudentForm().fields)
        self.assertIn('academic_period', StudentForm().fields)
        self.assertIn('course', [field.name for field in Subject._meta.fields])

    def test_attendance_models_and_teacher_summary_report_are_available(self):
        from django.apps import apps
        from .forms import ReportGenerationForm

        self.assertIs(apps.get_model('main_app', 'Attendance'), Attendance)
        self.assertIs(
            apps.get_model('main_app', 'AttendanceReport'), AttendanceReport)
        admin_types = {
            value for value, _ in ReportGenerationForm(role='admin').fields[
                'report_type'].choices
        }
        teacher_types = {
            value for value, _ in ReportGenerationForm(role='teacher').fields[
                'report_type'].choices
        }
        self.assertNotIn('attendance_summary', admin_types)
        self.assertIn('attendance_summary', teacher_types)

    def test_admin_can_view_learner_login_activity(self):
        admin = CustomUser.objects.create_user(
            email='activity-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='activity-learner@example.com',
            username='activity-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name='Activity',
            last_name='Learner',
        )
        course = Course.objects.create(name='Grade 7')
        learner.student.course = course
        learner.student.save()
        LoginActivity.objects.create(
            student=learner.student, ip_address='203.0.113.7')
        self.client.force_login(admin)

        response = self.client.get(reverse('learner_login_activity'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Activity Learner')
        self.assertContains(response, learner.username)
        self.assertContains(response, 'Grade 7')

    def test_learner_cannot_view_learner_login_activity(self):
        learner = CustomUser.objects.create_user(
            email='activity-viewer@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('learner_login_activity'))

        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

    def test_learner_login_activity_filters_by_name_username_date_and_course(self):
        admin = CustomUser.objects.create_user(
            email='filter-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        ada = CustomUser.objects.create_user(
            email='ada@example.com',
            username='ada-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name='Ada',
            last_name='Okafor',
        )
        binta = CustomUser.objects.create_user(
            email='binta@example.com',
            username='binta-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name='Binta',
            last_name='Bello',
        )
        grade_seven = Course.objects.create(name='Grade 7')
        grade_eight = Course.objects.create(name='Grade 8')
        ada.student.course = grade_seven
        ada.student.save()
        binta.student.course = grade_eight
        binta.student.save()
        LoginActivity.objects.create(student=ada.student)
        old_activity = LoginActivity.objects.create(student=binta.student)
        LoginActivity.objects.filter(pk=old_activity.pk).update(
            logged_in_at=timezone.now() - timedelta(days=1))
        self.client.force_login(admin)

        response = self.client.get(reverse('learner_login_activity'), {
            'learner': 'ada-learner',
            'date': timezone.localdate().isoformat(),
            'course': grade_seven.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, ada.username)
        self.assertNotContains(response, binta.username)
        self.assertContains(response, 'Grade 7')

        response = self.client.get(reverse('learner_login_activity'), {
            'learner': 'Binta',
        })
        self.assertContains(response, binta.username)
        self.assertNotContains(response, ada.username)

        response = self.client.get(reverse('learner_login_activity'), {
            'date': (timezone.localdate() - timedelta(days=1)).isoformat(),
        })
        self.assertContains(response, binta.username)
        self.assertNotContains(response, ada.username)

        response = self.client.get(reverse('learner_login_activity'), {
            'course': grade_seven.pk,
        })
        self.assertContains(response, ada.username)
        self.assertNotContains(response, binta.username)

    def test_admin_profile_cannot_change_password_but_remains_editable(self):
        admin = CustomUser.objects.create_user(
            email='profile-admin@example.com',
            username='profile-admin',
            password='original-password',
            user_type='1',
            gender='F',
            address='Initial address',
            first_name='Profile',
            last_name='Admin',
        )
        self.client.force_login(admin)

        with tempfile.TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                profile_url = reverse('admin_view_profile')
                first_update = self.client.post(profile_url, {
                    'first_name': 'Updated',
                    'last_name': 'Admin',
                    'current_password': 'original-password',
                    'new_password': 'new-unique-password-984!',
                    'confirm_password': 'new-unique-password-984!',
                    'profile_pic': self._profile_image_upload(),
                })

                self.assertRedirects(
                    first_update, profile_url, fetch_redirect_response=False)
                admin.refresh_from_db()
                self.assertTrue(admin.check_password('original-password'))
                self.assertFalse(
                    admin.check_password('new-unique-password-984!'))
                profile_response = self.client.get(profile_url)
                self.assertEqual(profile_response.status_code, 200)
                self.assertNotContains(profile_response, 'name="new_password"')
                self.assertNotContains(profile_response, 'name="reset_pin"')

                second_update = self.client.post(profile_url, {
                    'first_name': 'Updated Again',
                    'last_name': 'Admin',
                    'profile_pic': self._profile_image_upload(),
                })

        self.assertRedirects(
            second_update, profile_url, fetch_redirect_response=False)
        admin.refresh_from_db()
        self.assertEqual(admin.first_name, 'Updated Again')


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class TeacherRoleAndAttendanceTests(TestCase):
    def setUp(self):
        cache.clear()
        self.teacher = CustomUser.objects.create_user(
            email='teacher-role@example.com',
            username='teacher-role',
            password='valid-password',
            user_type='3',
            gender='F',
            address='',
            first_name='Taylor',
            last_name='Teacher',
        )
        self.admin = CustomUser.objects.create_user(
            email='teacher-role-admin@example.com',
            username='teacher-role-admin',
            password='valid-password',
            user_type='1',
            gender='M',
            address='',
        )
        self.course = None
        self.academic_period = None
        self.subject = None

    def _set_up_class(self):
        if self.course is not None:
            return
        self.course = Course.objects.create(name='Grade 8')
        self.academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01',
            end_year='2026-12-31',
        )
        self.subject = Subject.objects.create(
            name='Science',
            course=self.course,
            teacher=self.teacher.teacher,
        )

    def _learner(self, email, first_name):
        self._set_up_class()
        learner = CustomUser.objects.create_user(
            email=email,
            username=email.split('@', 1)[0],
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name=first_name,
            last_name='Learner',
        )
        learner.student.course = self.course
        learner.student.academic_period = self.academic_period
        learner.student.teacher = self.teacher.teacher
        learner.student.save()
        return learner

    @staticmethod
    def _image():
        image_data = BytesIO()
        Image.new('RGB', (1, 1), color='white').save(image_data, format='PNG')
        return SimpleUploadedFile(
            'teacher-profile.png',
            image_data.getvalue(),
            content_type='image/png',
        )

    def test_wrong_role_sign_in_is_rejected_for_every_role_pair(self):
        accounts = [
            ('admin', self.admin),
            ('teacher', self.teacher),
            ('learner', self._learner('role-pair@example.com', 'Role')),
        ]
        for account_role, account in accounts:
            for selected_role, _ in accounts:
                if account_role == selected_role:
                    continue
                with self.subTest(account=account_role, selected=selected_role):
                    self.client.logout()
                    response = self.client.post(reverse('user_login'), {
                        'username': account.username,
                        'password': 'valid-password',
                        'role': selected_role,
                    })
                    self.assertRedirects(
                        response,
                        reverse('role_login', args=[selected_role]),
                        fetch_redirect_response=False,
                    )
                    self.assertNotIn('_auth_user_id', self.client.session)

    def test_teacher_dashboard_and_profile_are_available(self):
        self.client.force_login(self.teacher)
        response = self.client.get(reverse('teacher_home'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Total Learners')
        self.assertContains(response, 'Total Subjects')

        response = self.client.post(reverse('teacher_profile'), {
            'first_name': 'Updated',
            'last_name': 'Teacher',
            'email': 'changed@example.com',
        })
        self.teacher.refresh_from_db()
        self.assertEqual(self.teacher.first_name, 'Updated')
        self.assertEqual(self.teacher.email, 'teacher-role@example.com')

    def test_teacher_can_manage_learners_and_admin_can_manage_teachers(self):
        self._set_up_class()
        self.client.force_login(self.teacher)
        response = self.client.get(reverse('manage_student'))
        self.assertEqual(response.status_code, 200)

        response = self.client.post(
            reverse('add_student'),
            {
                'first_name': 'Added',
                'last_name': 'Learner',
                'username': 'new-class-learner',
                'gender': 'F',
                'password': 'A-Unique-Strong-Password-289!',
                'address': 'School Road',
                'profile_pic': self._image(),
            },
        )
        self.assertRedirects(
            response, reverse('manage_student'),
            fetch_redirect_response=False)
        added_learner = CustomUser.objects.get(
            username='new-class-learner',
            user_type='2',
        )
        self.assertFalse(added_learner.pin_hash)

        learner = self._learner('teacher-edit-learner@example.com', 'Before')
        response = self.client.post(
            reverse('edit_student', args=[learner.student.pk]),
            {
                'first_name': 'After',
                'last_name': 'Learner',
                'username': 'updated-learner',
                'gender': 'F',
                'password': '',
                'profile_pic': self._image(),
                'address': 'School Road',
                'course': self.course.pk,
                'academic_period': self.academic_period.pk,
            },
        )
        self.assertEqual(
            response.status_code, 302,
            str(response.context['form'].errors) if response.context else '',
        )
        learner.refresh_from_db()
        self.assertEqual(learner.first_name, 'After')
        self.assertEqual(learner.username, 'updated-learner')
        response = self.client.post(
            reverse('delete_student', args=[learner.student.pk]))
        self.assertRedirects(
            response, reverse('manage_student'),
            fetch_redirect_response=False)
        self.assertFalse(CustomUser.objects.filter(pk=learner.pk).exists())

        self.client.force_login(self.admin)
        response = self.client.post(reverse('add_teacher'), {
            'first_name': 'New',
            'last_name': 'Teacher',
            'username': 'new-managed-teacher',
            'gender': 'M',
            'password': 'A-Unique-Strong-Password-289!',
        })
        self.assertEqual(response.status_code, 200)
        managed = CustomUser.objects.get(username='new-managed-teacher')
        self.assertTrue(check_password(
            response.context['issued_pin'], managed.pin_hash))
        self.assertFalse(managed.has_usable_password())
        teacher_profile = managed.teacher
        response = self.client.get(reverse('manage_teachers'))
        self.assertContains(response, managed.username)

        response = self.client.post(
            reverse('edit_teacher', args=[teacher_profile.pk]), {
                'first_name': 'Updated',
                'last_name': 'Teacher',
                'username': 'teacher-updated',
                'gender': 'F',
                'password': '',
            })
        self.assertRedirects(
            response, reverse('manage_teachers'),
            fetch_redirect_response=False)
        managed.refresh_from_db()
        self.assertEqual(managed.username, 'teacher-updated')

        response = self.client.post(
            reverse('delete_teacher', args=[teacher_profile.pk]))
        self.assertRedirects(
            response, reverse('manage_teachers'),
            fetch_redirect_response=False)
        self.assertFalse(CustomUser.objects.filter(pk=managed.pk).exists())

    def test_adding_learner_does_not_require_address_or_profile_picture(self):
        self.client.force_login(self.teacher)

        response = self.client.post(reverse('add_student'), {
            'first_name': 'No',
            'last_name': 'Photo',
            'username': 'no-photo-learner',
            'gender': 'F',
            'password': 'A-Unique-Strong-Password-289!',
        })

        self.assertRedirects(
            response, reverse('manage_student'),
            fetch_redirect_response=False)
        learner = CustomUser.objects.get(username='no-photo-learner')
        self.assertEqual(learner.address, '')
        self.assertFalse(learner.profile_pic)
        self.assertFalse(learner.pin_hash)
        self.assertEqual(learner.student.teacher, self.teacher.teacher)
        self.assertEqual(learner.student.created_by, self.teacher.teacher)

    def test_admin_and_teacher_pages_are_separated(self):
        self._set_up_class()
        learner = self._learner('restricted-learner@example.com', 'Private')
        result = StudentResult.objects.create(
            student=learner.student,
            subject=self.subject,
            test=20,
            exam=70,
        )
        admin_forbidden = (
            'manage_student_results',
            'generate_report_card',
        )
        self.client.force_login(self.admin)
        for page in admin_forbidden:
            with self.subTest(role='admin', page=page):
                response = self.client.get(reverse(page))
                self.assertRedirects(
                    response, reverse('admin_home'),
                    fetch_redirect_response=False)

        teacher_forbidden = (
            'manage_teachers',
            'add_teacher',
            'view_student_leave',
            'student_feedback_message',
            'admin_notify_student',
            'learner_login_activity',
            'admin_home',
            'admin_attendance_records',
        )
        self.client.force_login(self.teacher)
        for page in teacher_forbidden:
            with self.subTest(role='teacher', page=page):
                url = reverse(page)
                response = self.client.get(url)
                self.assertRedirects(
                    response, reverse('teacher_home'),
                    fetch_redirect_response=False)
        response = self.client.post(reverse('generate_report'), {
            'report_type': 'school_summary',
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotEqual(response.get('Content-Type'), 'application/pdf')

        self.client.force_login(learner)
        for page in (
                'teacher_home', 'teacher_attendance', 'manage_student',
                'manage_student_results', 'manage_teachers', 'admin_home'):
            with self.subTest(role='learner', page=page):
                response = self.client.get(reverse(page))
                self.assertRedirects(
                    response, reverse('student_home'),
                    fetch_redirect_response=False)
        self.assertTrue(result.pk)

    def test_teacher_attendance_is_corrected_not_duplicated_and_reported(self):
        self._set_up_class()
        learner = self._learner('attendance-learner@example.com', 'Attendee')
        self.client.force_login(self.teacher)
        selection = {
            'course': self.course.pk,
            'subject': self.subject.pk,
            'academic_period': self.academic_period.pk,
            'date': '2026-05-10',
            f'attendance_{learner.student.pk}': 'present',
        }
        response = self.client.post(reverse('teacher_attendance'), selection)
        self.assertRedirects(
            response,
            f"{reverse('teacher_attendance')}?course={self.course.pk}"
            f"&subject={self.subject.pk}&academic_period={self.academic_period.pk}&date=2026-05-10",
            fetch_redirect_response=False,
        )
        selection[f'attendance_{learner.student.pk}'] = 'absent'
        self.client.post(reverse('teacher_attendance'), selection)
        self.assertEqual(Attendance.objects.count(), 1)
        record = AttendanceReport.objects.get(
            student=learner.student,
            attendance__course=self.course,
            attendance__subject=self.subject,
            attendance__date='2026-05-10',
        )
        self.assertEqual(
            record.attendance.created_by, self.teacher.teacher)
        self.assertEqual(record.created_by, self.teacher.teacher)
        self.assertFalse(record.status)

        response = self.client.post(reverse('generate_report'), {
            'report_type': 'attendance_summary',
            'course': self.course.pk,
            'academic_period': self.academic_period.pk,
            'subject': self.subject.pk,
        })
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(b'/Subtype /Image', response.content)

        self.client.force_login(self.admin)
        response = self.client.get(reverse('admin_attendance_records'))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(reverse('teacher_attendance'), selection)
        self.assertRedirects(
            response, reverse('admin_home'),
            fetch_redirect_response=False)

    def test_learner_attendance_is_private_and_percentage_is_calculated(self):
        self._set_up_class()
        learner = self._learner('attendance-owner@example.com', 'Owner')
        other = self._learner('attendance-other@example.com', 'Other')
        attendance = Attendance.objects.create(
            course=self.course,
            subject=self.subject,
            academic_period=self.academic_period,
            date='2026-05-10',
        )
        AttendanceReport.objects.create(
            attendance=attendance, student=learner.student, status=True)
        AttendanceReport.objects.create(
            attendance=attendance, student=other.student, status=False)
        self.client.force_login(learner)

        response = self.client.get(reverse('student_attendance'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '100.0%')
        self.assertContains(response, 'Present')
        self.assertNotContains(
            response, '<span class="badge badge-danger">Absent</span>')

    def test_admin_can_assign_learners_and_subjects_to_teachers(self):
        self._set_up_class()
        learner = CustomUser.objects.create_user(
            username='unassigned-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        subject = Subject.objects.create(name='Unassigned Subject')
        self.client.force_login(self.admin)

        response = self.client.post(reverse('manage_assignments'), {
            'assignment_type': 'learners',
            'learners-learners': [learner.student.pk],
            'learners-teacher': self.teacher.teacher.pk,
        })
        self.assertRedirects(
            response, reverse('manage_assignments'),
            fetch_redirect_response=False)
        learner.student.refresh_from_db()
        self.assertEqual(learner.student.teacher, self.teacher.teacher)
        self.assertIsNone(learner.student.created_by)

        response = self.client.post(reverse('manage_assignments'), {
            'assignment_type': 'subjects',
            'subjects-subjects': [subject.pk],
            'subjects-teacher': self.teacher.teacher.pk,
        })
        self.assertRedirects(
            response, reverse('manage_assignments'),
            fetch_redirect_response=False)
        subject.refresh_from_db()
        self.assertEqual(subject.teacher, self.teacher.teacher)
        self.assertIsNone(subject.created_by)

        self.client.force_login(self.teacher)
        response = self.client.post(reverse('add_subject'), {
            'name': 'Teacher Created Subject',
            'course': self.course.pk,
        })
        self.assertRedirects(
            response, reverse('manage_subject'),
            fetch_redirect_response=False)
        created_subject = Subject.objects.get(name='Teacher Created Subject')
        self.assertEqual(created_subject.teacher, self.teacher.teacher)
        self.assertEqual(created_subject.created_by, self.teacher.teacher)

    def test_teacher_cannot_access_another_teachers_learner_or_records(self):
        self._set_up_class()
        other_teacher = CustomUser.objects.create_user(
            username='other-teacher',
            password='valid-password',
            user_type='3',
            gender='M',
        )
        learner = CustomUser.objects.create_user(
            username='other-assigned-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        learner.student.course = self.course
        learner.student.academic_period = self.academic_period
        learner.student.teacher = other_teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(
            name='Other Teacher Subject',
            course=self.course,
            teacher=other_teacher.teacher,
        )
        result = StudentResult.objects.create(
            student=learner.student, subject=subject, test=10, exam=20)
        self.client.force_login(self.teacher)

        response = self.client.get(reverse('manage_student'))
        self.assertNotContains(response, learner.username)
        self.assertEqual(response.status_code, 200)
        for url_name, args in (
            ('edit_student', [learner.student.pk]),
            ('edit_student_result', [result.pk]),
            ('delete_student_result', [result.pk]),
            ('edit_subject', [subject.pk]),
            ('delete_subject', [subject.pk]),
        ):
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name, args=args))
                self.assertEqual(response.status_code, 404)

        response = self.client.post(reverse('generate_report_card'), {
            'student': learner.student.pk,
            'academic_period': self.academic_period.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotEqual(response.get('Content-Type'), 'application/pdf')

    def test_deleting_teacher_preserves_learners_subjects_and_academic_records(self):
        self._set_up_class()
        learner = self._learner('preserved-learner@example.com', 'Preserved')
        learner.student.created_by = self.teacher.teacher
        learner.student.save()
        result = StudentResult.objects.create(
            student=learner.student,
            subject=self.subject,
            created_by=self.teacher.teacher,
            test=20,
            exam=70,
        )
        attendance = Attendance.objects.create(
            course=self.course,
            subject=self.subject,
            academic_period=self.academic_period,
            created_by=self.teacher.teacher,
            date='2026-05-10',
        )
        attendance_report = AttendanceReport.objects.create(
            attendance=attendance,
            student=learner.student,
            created_by=self.teacher.teacher,
            status=True,
        )

        self.teacher.delete()

        learner.student.refresh_from_db()
        self.subject.refresh_from_db()
        result.refresh_from_db()
        attendance.refresh_from_db()
        attendance_report.refresh_from_db()
        self.assertIsNone(learner.student.teacher)
        self.assertIsNone(learner.student.created_by)
        self.assertTrue(Student.objects.filter(pk=learner.student.pk).exists())
        self.assertIsNone(self.subject.teacher)
        self.assertIsNone(self.subject.created_by)
        self.assertTrue(Subject.objects.filter(pk=self.subject.pk).exists())
        self.assertIsNone(result.created_by)
        self.assertIsNone(attendance.created_by)
        self.assertIsNone(attendance_report.created_by)

    def test_grade_course_order_is_numeric(self):
        names = [f'Grade {number}' for number in range(1, 11)]
        for name in reversed(names):
            Course.objects.create(name=name)

        from .models import grade_ordered_courses

        self.assertEqual(
            list(grade_ordered_courses().values_list('name', flat=True)),
            names,
        )

    @staticmethod
    def _profile_image_upload():
        image_data = BytesIO()
        Image.new('RGB', (1, 1), color='white').save(image_data, format='PNG')
        return SimpleUploadedFile(
            'profile.png', image_data.getvalue(), content_type='image/png')

    def test_admin_can_download_pdf_reports(self):
        admin = CustomUser.objects.create_user(
            email='report-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
            first_name='Report',
            last_name='Admin',
        )
        learner = CustomUser.objects.create_user(
            email='report-learner@example.com',
            username='report-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name='Test',
            last_name='Learner',
        )
        course = Course.objects.create(name='Grade 7')
        academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        subject = Subject.objects.create(name='Mathematics', course=course)
        learner.student.course = course
        learner.student.academic_period = academic_period
        learner.student.save()
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=18, exam=72)
        self.client.force_login(admin)

        report_options = [
            {'report_type': 'school_summary'},
            {
                'report_type': 'class_list',
                'course': course.pk,
                'academic_period': academic_period.pk,
            },
        ]
        for options in report_options:
            with self.subTest(report_type=options['report_type']):
                response = self.client.post(reverse('generate_report'), options)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response['Content-Type'], 'application/pdf')
                self.assertTrue(response.content.startswith(b'%PDF'))
                self.assertIn(b'/Subtype /Image', response.content)
                self.assertIn('attachment;', response['Content-Disposition'])

        from .hod_views import _report_data

        _, headers, rows = _report_data(
            'class_list', course, academic_period, None)
        self.assertEqual(headers, ['Learner name', 'Username'])
        self.assertEqual(rows[0][1], learner.username)

    def test_teacher_can_download_results_report(self):
        teacher = CustomUser.objects.create_user(
            email='report-teacher@example.com',
            password='valid-password',
            user_type='3',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='report-result-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        course = Course.objects.create(name='Grade 7')
        academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        subject = Subject.objects.create(name='Mathematics', course=course)
        learner.student.teacher = teacher.teacher
        learner.student.course = course
        learner.student.academic_period = academic_period
        learner.student.save()
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=18, exam=72)
        self.client.force_login(teacher)

        response = self.client.post(reverse('generate_report'), {
            'report_type': 'learner_results',
            'course': course.pk,
            'academic_period': academic_period.pk,
            'subject': subject.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn(b'/Subtype /Image', response.content)

    def test_teacher_can_add_and_update_result_and_learner_sees_latest_scores(self):
        teacher = CustomUser.objects.create_user(
            email='result-manager@example.com',
            password='valid-password',
            user_type='3',
            gender='F',
            address='',
        )
        learner = CustomUser.objects.create_user(
            email='result-owner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name='Alex',
            last_name='Learner',
        )
        learner.student.teacher = teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(name='Mathematics')
        self.client.force_login(teacher)

        response = self.client.post(reverse('manage_student_results'), {
            'student': learner.student.pk,
            'subject': subject.pk,
            'test': '18',
            'exam': '72',
        })

        self.assertRedirects(
            response,
            reverse('manage_student_results'),
            fetch_redirect_response=False,
        )
        result = StudentResult.objects.get(
            student=learner.student, subject=subject)
        self.assertEqual(result.test, 18)
        self.assertEqual(result.exam, 72)
        self.assertEqual(result.created_by, teacher.teacher)

        response = self.client.post(
            reverse('edit_student_result', args=[result.pk]),
            {
                'student': learner.student.pk,
                'subject': subject.pk,
                'test': '20',
                'exam': '75',
            },
        )
        self.assertRedirects(
            response,
            reverse('manage_student_results'),
            fetch_redirect_response=False,
        )

        self.client.force_login(learner)
        response = self.client.get(reverse('student_view_result'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Mathematics')
        self.assertContains(response, '20.0')
        self.assertContains(response, '75.0')
        self.assertContains(response, '95.0')

    def test_learner_cannot_access_admin_result_management(self):
        learner = CustomUser.objects.create_user(
            email='result-access-denied@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('manage_student_results'))

        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

    def test_learner_is_redirected_from_report_page(self):
        learner = CustomUser.objects.create_user(
            email='report-denied@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('generate_report'))

        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

    def test_empty_report_shows_message_instead_of_returning_pdf(self):
        admin = CustomUser.objects.create_user(
            email='empty-report-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.post(reverse('generate_report'), {
            'report_type': 'school_summary',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/html; charset=utf-8')
        self.assertContains(
            response, 'There is no data for the selected report options.')

    def test_admin_can_sign_in_through_admin_interface(self):
        admin = CustomUser.objects.create_user(
            username='admin-login-interface',
            password='valid-password',
            user_type=1,
            gender='F',
            address='',
        )
        admin_pin = issue_pin(admin)

        response = self.client.post(reverse('user_login'), {
            'username': admin.username,
            'pin': admin_pin,
            'role': 'admin',
        })

        self.assertRedirects(
            response,
            reverse('admin_home'),
            fetch_redirect_response=False,
        )

    def test_admin_cannot_sign_in_through_learner_interface(self):
        CustomUser.objects.create_user(
            email='admin@example.com',
            password='valid-password',
            user_type=1,
            gender='F',
            address='',
        )

        response = self.client.post(reverse('user_login'), {
            'username': 'admin@example.com',
            'password': 'valid-password',
            'role': 'learner',
        })

        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_wrong_password_does_not_sign_in(self):
        CustomUser.objects.create_user(
            email='learner@example.com',
            password='valid-password',
            user_type=2,
            gender='F',
            address='',
        )

        response = self.client.post(reverse('user_login'), {
            'username': 'learner@example.com',
            'password': 'incorrect-password',
            'role': 'learner',
        })

        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)


class GradeBandAndReportCardTests(TestCase):
    def setUp(self):
        self.teacher = CustomUser.objects.create_user(
            email='grade-teacher@example.com',
            password='valid-password',
            user_type='3',
            gender='F',
            address='',
        )
        self.client.force_login(self.teacher)

    @staticmethod
    def _grade_band_payload(rows):
        data = {
            'action': 'save_grade_bands',
            'bands-TOTAL_FORMS': '8',
            'bands-INITIAL_FORMS': '0',
            'bands-MIN_NUM_FORMS': '0',
            'bands-MAX_NUM_FORMS': '8',
        }
        for index in range(8):
            row = rows[index] if index < len(rows) else {}
            for name in ('from_score', 'to_score', 'symbol'):
                data[f'bands-{index}-{name}'] = row.get(name, '')
        return data

    def test_reports_page_has_eight_grade_band_rows_and_saves_valid_rows(self):
        GradeBand.objects.create(
            symbol='Old', minimum_score=0, maximum_score=100)
        response = self.client.get(reverse('generate_report'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Grade Bands')
        self.assertEqual(len(response.context['grade_band_formset'].forms), 8)
        self.assertContains(response, f'Scores are out of {settings.SIMS_MAX_TOTAL_SCORE}.')

        response = self.client.post(
            reverse('generate_report'),
            self._grade_band_payload([
                {'from_score': '0', 'to_score': '69', 'symbol': 'C'},
                {'from_score': '70', 'to_score': '79', 'symbol': 'B'},
                {
                    'from_score': '80',
                    'to_score': str(settings.SIMS_MAX_TOTAL_SCORE),
                    'symbol': 'A',
                },
            ]),
        )

        self.assertRedirects(
            response, reverse('generate_report'),
            fetch_redirect_response=False)
        self.assertEqual(
            list(GradeBand.objects.order_by('minimum_score').values_list(
                'symbol', flat=True)),
            ['C', 'B', 'A'],
        )
        self.assertEqual(
            GradeBand.objects.get(symbol='A').description, '')

    def test_grade_band_rows_reject_incomplete_inverted_overlapping_and_duplicate(self):
        existing = GradeBand.objects.create(
            symbol='Keep', minimum_score=0, maximum_score=10)
        invalid_rows = (
            (
                {'from_score': '1', 'to_score': '2', 'symbol': ''},
                'Complete all three fields',
            ),
            (
                {'from_score': '80', 'to_score': '70', 'symbol': 'A'},
                'greater than or equal to From',
            ),
            (
                {'from_score': '0', 'to_score': '70', 'symbol': 'A'},
                'overlaps row',
            ),
            (
                {'from_score': '0', 'to_score': '40', 'symbol': 'A'},
                'duplicated',
            ),
        )
        for first, error in invalid_rows:
            with self.subTest(error=error):
                second = (
                    {'from_score': '60', 'to_score': '80', 'symbol': 'B'}
                    if 'overlaps' in error
                    else {'from_score': '50', 'to_score': '70', 'symbol': 'A'}
                    if 'duplicated' in error
                    else {}
                )
                response = self.client.post(
                    reverse('generate_report'),
                    self._grade_band_payload([first, second]),
                )
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, error)
                self.assertTrue(GradeBand.objects.filter(pk=existing.pk).exists())

    def test_grade_bands_above_configured_maximum_are_rejected(self):
        response = self.client.post(
            reverse('generate_report'),
            self._grade_band_payload([{
                'from_score': '0',
                'to_score': str(settings.SIMS_MAX_TOTAL_SCORE + 1),
                'symbol': 'A',
            }]),
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Ensure this value is less than or equal to')
        self.assertFalse(GradeBand.objects.exists())

    def test_grade_band_gaps_warn_but_do_not_prevent_saving(self):
        response = self.client.post(
            reverse('generate_report'),
            self._grade_band_payload([
                {'from_score': '0', 'to_score': '49', 'symbol': 'C'},
                {'from_score': '60', 'to_score': '79', 'symbol': 'B'},
            ]),
            follow=True,
        )
        self.assertContains(response, 'Some score ranges are not covered')
        self.assertEqual(GradeBand.objects.count(), 2)

    def test_grade_helper_rounds_half_up_and_rejects_totals_outside_limit(self):
        from .grading import get_grade_band

        lower = GradeBand.objects.create(
            symbol='B', minimum_score=70, maximum_score=79)
        upper = GradeBand.objects.create(
            symbol='A', minimum_score=80,
            maximum_score=settings.SIMS_MAX_TOTAL_SCORE)

        self.assertEqual(get_grade_band(79.5).pk, upper.pk)
        self.assertEqual(get_grade_band(79.4).pk, lower.pk)
        self.assertEqual(get_grade_band(-0.1), 'Not graded')
        self.assertEqual(
            get_grade_band(settings.SIMS_MAX_TOTAL_SCORE + 0.1),
            'Not graded',
        )

    def test_changing_grade_bands_changes_grade_for_existing_result(self):
        learner = CustomUser.objects.create_user(
            email='grade-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        learner.student.teacher = self.teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(name='Mathematics')
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=20, exam=70)
        GradeBand.objects.create(
            symbol='B', minimum_score=70, maximum_score=89)

        response = self.client.get(reverse('manage_student_results'))
        self.assertContains(response, '90.0 (Not graded)')

        response = self.client.post(reverse('generate_report'),
            self._grade_band_payload([{
                'from_score': '0', 'to_score': '89', 'symbol': 'B',
            }, {
                'from_score': '90',
                'to_score': str(settings.SIMS_MAX_TOTAL_SCORE),
                'symbol': 'A',
            }]))
        self.assertRedirects(
            response,
            reverse('generate_report'),
            fetch_redirect_response=False,
        )
        response = self.client.get(reverse('manage_student_results'))
        self.assertContains(response, '90.0 (A)')

    def test_result_entry_rejects_scores_outside_configured_total(self):
        learner = CustomUser.objects.create_user(
            username='score-limit-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        learner.student.teacher = self.teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(name='Mathematics')
        existing = StudentResult.objects.create(
            student=learner.student, subject=subject, test=20, exam=30)

        response = self.client.post(reverse('manage_student_results'), {
            'student': learner.student.pk,
            'subject': subject.pk,
            'test': settings.SIMS_MAX_TOTAL_SCORE,
            'exam': 1,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'cannot total more than')
        self.assertEqual(StudentResult.objects.count(), 1)

        response = self.client.post(
            reverse('edit_student_result', args=[existing.pk]), {
                'student': learner.student.pk,
                'subject': subject.pk,
                'test': settings.SIMS_MAX_TOTAL_SCORE,
                'exam': 1,
            })
        self.assertEqual(response.status_code, 200)
        existing.refresh_from_db()
        self.assertEqual((existing.test, existing.exam), (20, 30))

        response = self.client.post(reverse('manage_student_results'), {
            'student': learner.student.pk,
            'subject': subject.pk,
            'test': -1,
            'exam': 0,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['form'].is_valid())
        self.assertEqual(StudentResult.objects.count(), 1)

    @override_settings(SIMS_MAX_TOTAL_SCORE=50)
    def test_configured_maximum_drives_grade_band_and_result_validation(self):
        response = self.client.get(reverse('generate_report'))
        self.assertContains(response, 'Scores are out of 50.')

        response = self.client.post(
            reverse('generate_report'),
            self._grade_band_payload([{
                'from_score': '0', 'to_score': '51', 'symbol': 'A',
            }]),
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(GradeBand.objects.exists())

        learner = CustomUser.objects.create_user(
            username='setting-limit-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        learner.student.teacher = self.teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(name='Science')
        response = self.client.post(reverse('manage_student_results'), {
            'student': learner.student.pk,
            'subject': subject.pk,
            'test': 30,
            'exam': 21,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(StudentResult.objects.exists())

    def test_check_result_totals_lists_invalid_rows_without_modifying_them(self):
        learner = CustomUser.objects.create_user(
            username='legacy-score-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        subject = Subject.objects.create(name='Legacy Scores')
        over_limit = StudentResult.objects.create(
            student=learner.student,
            subject=subject,
            test=settings.SIMS_MAX_TOTAL_SCORE + 5,
            exam=1,
        )
        below_zero = StudentResult.objects.create(
            student=learner.student, subject=subject, test=-2, exam=1)
        output = StringIO()

        call_command('check_result_totals', stdout=output)

        self.assertIn(f'Result {over_limit.pk}', output.getvalue())
        self.assertIn('above maximum', output.getvalue())
        self.assertIn(f'Result {below_zero.pk}', output.getvalue())
        self.assertIn('below zero', output.getvalue())
        over_limit.refresh_from_db()
        below_zero.refresh_from_db()
        self.assertEqual(
            (over_limit.test, over_limit.exam),
            (settings.SIMS_MAX_TOTAL_SCORE + 5, 1),
        )
        self.assertEqual((below_zero.test, below_zero.exam), (-2, 1))

    def test_learner_results_show_their_current_grade(self):
        learner = CustomUser.objects.create_user(
            email='learner-grade@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        subject = Subject.objects.create(name='Science')
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=30, exam=50)
        GradeBand.objects.create(
            symbol='B', minimum_score=80, maximum_score=80,
            description='Good')
        self.client.force_login(learner)

        response = self.client.get(reverse('student_view_result'))

        self.assertContains(response, '80.0 (B)')

    def test_removed_grade_band_urls_return_404_and_reports_respect_roles(self):
        for url in (
            '/admin/grade-bands/',
            '/admin/grade-bands/edit/1/',
            '/admin/grade-bands/delete/1/',
        ):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 404)

        learner = CustomUser.objects.create_user(
            email='grade-viewer@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        self.client.force_login(learner)

        response = self.client.get(reverse('generate_report'))
        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

        self.client.logout()
        response = self.client.get('/admin/grade-bands/')
        self.assertRedirects(
            response,
            reverse('login_page'),
            fetch_redirect_response=False,
        )
        admin = CustomUser.objects.create_user(
            email='grade-admin-restricted@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)
        response = self.client.get(reverse('generate_report'))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Grade Bands')

    def test_report_card_pdf_downloads_for_teacher(self):
        learner, academic_period, _, _ = self._make_learner_with_result(
            'card-learner@example.com', 'Grade 8', 'Biology')
        GradeBand.objects.create(
            symbol='A', minimum_score=80, maximum_score=100,
            description='Excellent', remark='Outstanding')

        response = self.client.post(reverse('generate_report_card'), {
            'student': learner.student.pk,
            'academic_period': academic_period.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(b'/Subtype /Image', response.content)
        self.assertIn(
            'attachment; filename="learner-report-card-from-2026-01-01-to-2026-12-31.pdf"',
            response['Content-Disposition'],
        )

    def test_teacher_report_card_shows_message_when_there_are_no_results(self):
        learner = CustomUser.objects.create_user(
            email='no-card-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        learner.student.teacher = self.teacher.teacher
        learner.student.academic_period = academic_period
        learner.student.save()

        response = self.client.post(reverse('generate_report_card'), {
            'student': learner.student.pk,
            'academic_period': academic_period.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/html; charset=utf-8')
        self.assertContains(response, 'There are no results for this learner and academic period.')

    def test_learner_can_download_only_their_own_report_card(self):
        learner, academic_period, own_subject, _ = self._make_learner_with_result(
            'own-card@example.com', 'Grade 7', 'Own Subject')
        other_learner, _, other_subject, _ = self._make_learner_with_result(
            'other-card@example.com', 'Grade 8', 'Other Subject')
        self.client.force_login(learner)

        with patch(
                'main_app.hod_views._create_report_pdf',
                return_value=HttpResponse(b'%PDF-own-only')) as pdf_builder:
            response = self.client.get(
                reverse('student_report_card'),
                {'student': other_learner.student.pk},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b'%PDF-own-only')
        report_rows = pdf_builder.call_args.args[3]
        self.assertEqual(report_rows[0][0], 'Own Subject')
        self.assertNotIn(other_subject.name, [
            value for row in report_rows for value in row])
        self.assertNotEqual(own_subject.pk, other_subject.pk)

        response = self.client.get(reverse('generate_report_card'))
        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

    def test_learner_cannot_view_another_learners_report_card(self):
        learner, _, _, _ = self._make_learner_with_result(
            'own-private-card@example.com', 'Grade 7', 'Private Subject')
        other_learner, _, _, _ = self._make_learner_with_result(
            'other-private-card@example.com', 'Grade 8', 'Other Private Subject')
        self.client.force_login(learner)

        response = self.client.get(
            reverse('student_report_card'),
            {'learner': other_learner.student.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def _make_learner_with_result(self, email, course_name, subject_name):
        learner = CustomUser.objects.create_user(
            email=email,
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
            first_name=email.split('@')[0],
            last_name='Learner',
        )
        course = Course.objects.create(name=course_name)
        academic_period = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        learner.student.course = course
        learner.student.academic_period = academic_period
        learner.student.teacher = self.teacher.teacher
        learner.student.save()
        subject = Subject.objects.create(
            name=subject_name, course=course, teacher=self.teacher.teacher)
        result = StudentResult.objects.create(
            student=learner.student, subject=subject, test=20, exam=70)
        return learner, academic_period, subject, result


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class RegistrationAndRestrictedProfileTests(TestCase):
    def setUp(self):
        cache.clear()

    def _learner_registration_data(self, username='new-learner', **kwargs):
        data = {
            'first_name': 'Taylor',
            'last_name': 'Learner',
            'username': username,
            'gender': 'F',
            'password': 'A-Unique-Strong-Password-289!',
            'confirm_password': 'A-Unique-Strong-Password-289!',
        }
        data.update(kwargs)
        return data

    def test_learner_registration_then_learner_sign_in(self):
        data = self._learner_registration_data()

        response = self.client.post(
            reverse('register_learner'), data, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Registration complete')
        learner = CustomUser.objects.get(username='new-learner')
        self.assertEqual(learner.user_type, '2')
        self.assertIsNone(learner.student.course)
        self.assertIsNone(learner.student.academic_period)
        self.assertIsNone(learner.student.teacher)
        self.assertIsNone(learner.student.created_by)
        self.assertFalse(learner.pin_hash)
        self.assertFalse(learner.profile_pic)
        self.assertTrue(learner.check_password(data['password']))
        self.assertFalse(learner.email)
        self.assertNotIn('_auth_user_id', self.client.session)

        response = self.client.post(reverse('user_login'), {
            'username': learner.username,
            'password': data['password'],
            'role': 'learner',
        })
        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )

    def test_duplicate_learner_username_is_rejected_case_insensitively(self):
        CustomUser.objects.create_user(
            username='duplicate-user',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        learner_data = self._learner_registration_data(
            username='DUPLICATE-USER')
        response = self.client.post(
            reverse('register_learner'), learner_data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'already exists')
        self.assertEqual(CustomUser.objects.filter(
            username__iexact='duplicate-user').count(), 1)

    def test_password_confirmation_and_minimum_length_validation(self):
        data = self._learner_registration_data(
            confirm_password='A-different-password-289!')
        response = self.client.post(
            reverse('register_learner'), data)
        self.assertContains(response, 'Passwords do not match')

        data = self._learner_registration_data(
            username='weak-password',
            password='short',
            confirm_password='short',
        )
        response = self.client.post(
            reverse('register_learner'), data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response, 'Password must be at least 8 characters long.')
        self.assertFalse(CustomUser.objects.filter(
            username='weak-password').exists())

    def test_eight_character_password_accepts_any_character_combination(self):
        data = self._learner_registration_data(
            username='numeric-password',
            password='12345678',
            confirm_password='12345678',
        )
        response = self.client.post(
            reverse('register_learner'), data)

        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
        )
        learner = CustomUser.objects.get(username='numeric-password')
        self.assertTrue(learner.check_password('12345678'))

    def test_registration_preserves_spaces_in_passwords(self):
        password = '  12345 '
        data = self._learner_registration_data(
            username='spaced-password',
            password=password,
            confirm_password=password,
        )
        response = self.client.post(
            reverse('register_learner'), data)

        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
        )
        learner = CustomUser.objects.get(username='spaced-password')
        self.assertTrue(learner.check_password(password))

    def test_registration_page_shows_prominent_crest_and_live_password_feedback(self):
        response = self.client.get(reverse('register_learner'))

        self.assertContains(response, 'C. Heuva school crest')
        self.assertContains(
            response, 'main_app/js/registration-validation.js')
        self.assertContains(response, 'password-length-feedback')
        self.assertContains(response, 'password-match-feedback')
        self.assertContains(
            response, 'At least 8 characters. Any characters are allowed.')

    def test_learner_registration_does_not_show_or_require_a_grade(self):
        Course.objects.create(name='Grade 9')
        response = self.client.get(reverse('register_learner'))
        self.assertNotContains(response, 'Grade 9')
        self.assertNotContains(response, 'name="course"')

        data = self._learner_registration_data(username='unassigned-learner')
        response = self.client.post(
            reverse('register_learner'), data)
        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
        )
        self.assertIsNone(
            CustomUser.objects.get(username='unassigned-learner').student.course)

    def test_invalid_usernames_are_rejected(self):
        Course.objects.create(name='Grade 7')
        for username in ('abc', 'a' * 21, '_starts-wrong', 'has space'):
            with self.subTest(username=username):
                data = self._learner_registration_data(username=username)
                response = self.client.post(
                    reverse('register_learner'), data)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['form'].errors.get('username'))
                self.assertFalse(
                    CustomUser.objects.filter(username__iexact=username).exists())

    def test_username_command_creates_unique_names_and_csv(self):
        first = CustomUser.objects.create_user(
            email='ab@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        second = CustomUser.objects.create_user(
            email='AB@another.example',
            password='valid-password',
            user_type='3',
            gender='M',
        )

        with tempfile.TemporaryDirectory() as directory:
            csv_path = f'{directory}\\usernames.csv'
            call_command('create_usernames', csv_path, stdout=StringIO())

            first.refresh_from_db()
            second.refresh_from_db()
            self.assertEqual(first.username, 'ab00')
            self.assertEqual(second.username, 'ab001')
            with open(csv_path, newline='', encoding='utf-8') as csv_file:
                rows = csv_file.read().splitlines()
            self.assertEqual(rows[0], 'name,role,username')
            self.assertEqual(len(rows), 3)

    def test_create_admin_command_creates_account_and_displays_pin_once(self):
        with patch(
            'builtins.input',
            side_effect=['Alex', 'Administrator', 'new-admin', 'F'],
        ):
            output = StringIO()
            call_command('create_admin', stdout=output)

        admin = CustomUser.objects.get(username='new-admin')
        self.assertEqual(admin.user_type, '1')
        self.assertEqual(admin.first_name, 'Alex')
        self.assertEqual(admin.last_name, 'Administrator')
        self.assertTrue(Admin.objects.filter(admin=admin).exists())
        self.assertFalse(admin.has_usable_password())
        self.assertTrue(check_password(
            output.getvalue().split('One-time PIN: ', 1)[1].strip(),
            admin.pin_hash,
        ))
        self.assertEqual(output.getvalue().count('One-time PIN:'), 1)

    def test_legacy_email_sign_in_can_be_disabled(self):
        account = CustomUser.objects.create_user(
            email='legacy@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        self.assertIsNone(account.username)

        with override_settings(SIMS_ALLOW_EMAIL_SIGNIN=True):
            response = self.client.post(reverse('user_login'), {
                'username': 'legacy@example.com',
                'password': 'valid-password',
                'role': 'learner',
            })
        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )
        self.client.logout()

        with override_settings(SIMS_ALLOW_EMAIL_SIGNIN=False):
            response = self.client.post(reverse('user_login'), {
                'username': 'legacy@example.com',
                'password': 'valid-password',
                'role': 'learner',
            })
        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_learner_registration_is_available_without_database_grades(self):
        response = self.client.get(reverse('register_learner'))

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Registration is not available')
        self.assertNotContains(response, ' disabled')

    def test_registration_rejects_invalid_profile_image(self):
        data = self._learner_registration_data(username='invalid-image')
        data['profile_pic'] = SimpleUploadedFile(
            'photo.jpg', b'this is not an image', content_type='image/jpeg')

        response = self.client.post(
            reverse('register_learner'),
            data,
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Upload a valid image')
        self.assertFalse(CustomUser.objects.filter(
            username='invalid-image').exists())

    def test_staff_registration_urls_return_404(self):
        for role in ('admin', 'teacher'):
            with self.subTest(role=role):
                response = self.client.get(f'/register/{role}/')
                self.assertEqual(response.status_code, 404)

    def test_staff_sign_in_pages_do_not_link_to_registration(self):
        for role in ('admin', 'teacher'):
            with self.subTest(role=role):
                response = self.client.get(
                    reverse('role_login', args=[role]))
                self.assertEqual(response.status_code, 200)
                self.assertNotContains(response, 'Register')
                self.assertNotContains(response, '/register/')

    def test_learner_sign_in_page_links_to_learner_registration(self):
        response = self.client.get(
            reverse('role_login', args=['learner']))
        self.assertContains(response, "Don't have an account? Register")
        self.assertContains(response, reverse('register_learner'))

        response = self.client.get(reverse('register_learner'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            reverse('role_login', args=['learner']),
        )

    def test_signed_in_user_is_redirected_from_registration(self):
        admin = CustomUser.objects.create_user(
            email='already-signed-in@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.get(reverse('register_learner'))

        self.assertRedirects(
            response,
            reverse('admin_home'),
            fetch_redirect_response=False,
        )

    def test_registered_roles_cannot_sign_in_through_the_other_role(self):
        learner_data = self._learner_registration_data(
            username='role-check-learner')
        learner_registration = self.client.post(
            reverse('register_learner'), learner_data)
        response = self.client.post(reverse('user_login'), {
            'username': learner_data['username'],
            'password': learner_data['password'],
            'role': 'admin',
        })
        self.assertRedirects(
            response,
            reverse('role_login', args=['admin']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

        admin = CustomUser.objects.create_user(
            username='role-check-admin',
            password='A-Unique-Strong-Password-289!',
            user_type='1',
            gender='M',
        )
        admin_pin = issue_pin(admin)
        response = self.client.post(reverse('user_login'), {
            'username': admin.username,
            'password': 'A-Unique-Strong-Password-289!',
            'pin': admin_pin,
            'role': 'learner',
        })
        self.assertRedirects(
            response,
            reverse('role_login', args=['learner']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_registration_attempts_are_rate_limited_per_client(self):
        from django.test import RequestFactory
        from .views import _registration_rate_limited

        request = RequestFactory().post(
            '/register/learner/', REMOTE_ADDR='192.0.2.55')
        self.assertEqual(
            [_registration_rate_limited(request) for _ in range(6)],
            [False, False, False, False, False, True],
        )

    def test_learner_profile_ignores_submitted_email_grade_and_session(self):
        learner = CustomUser.objects.create_user(
            email='restricted-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        own_grade = Course.objects.create(name='Own Grade')
        forged_grade = Course.objects.create(name='Forged Grade')
        own_session = AcademicPeriod.objects.create(
            start_year='2026-01-01', end_year='2026-12-31')
        forged_session = AcademicPeriod.objects.create(
            start_year='2027-01-01', end_year='2027-12-31')
        learner.student.course = own_grade
        learner.student.academic_period = own_session
        learner.student.save()
        self.client.force_login(learner)

        response = self.client.post(reverse('student_view_profile'), {
            'first_name': 'Changed',
            'last_name': 'Name',
            'email': 'attacker@example.com',
            'course': forged_grade.pk,
            'academic_period': forged_session.pk,
        })

        self.assertRedirects(
            response,
            reverse('student_view_profile'),
            fetch_redirect_response=False,
        )
        learner.refresh_from_db()
        learner.student.refresh_from_db()
        self.assertEqual(learner.email, 'restricted-learner@example.com')
        self.assertEqual(learner.student.course, own_grade)
        self.assertEqual(learner.student.academic_period, own_session)
        self.assertEqual(learner.first_name, 'Changed')

    def test_admin_profile_ignores_submitted_email(self):
        admin = CustomUser.objects.create_user(
            email='restricted-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(admin)

        response = self.client.post(reverse('admin_view_profile'), {
            'first_name': 'Changed',
            'last_name': 'Admin',
            'email': 'attacker@example.com',
        })

        self.assertRedirects(
            response,
            reverse('admin_view_profile'),
            fetch_redirect_response=False,
        )
        admin.refresh_from_db()
        self.assertEqual(admin.email, 'restricted-admin@example.com')

    def test_profile_pages_are_limited_to_the_signed_in_users_role(self):
        learner = CustomUser.objects.create_user(
            email='profile-owner-learner@example.com',
            username='profile-owner-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        admin = CustomUser.objects.create_user(
            email='profile-owner-admin@example.com',
            username='profile-owner-admin',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        self.client.force_login(learner)
        response = self.client.get(reverse('admin_view_profile'))
        self.assertRedirects(
            response,
            reverse('student_home'),
            fetch_redirect_response=False,
        )
        response = self.client.get(reverse('student_view_profile'))
        self.assertEqual(response.status_code, 200)

        self.client.force_login(admin)
        response = self.client.get(reverse('student_view_profile'))
        self.assertRedirects(
            response,
            reverse('admin_home'),
            fetch_redirect_response=False,
        )
        response = self.client.get(reverse('admin_view_profile'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'My Profile')
        self.assertContains(response, admin.username)
        self.assertNotContains(response, 'name="email"')

    def test_profile_password_change_requires_current_password_and_keeps_session(self):
        learner = CustomUser.objects.create_user(
            email='password-profile@example.com',
            password='Original-Strong-Password-789!',
            user_type='2',
            gender='F',
            address='',
            first_name='Password',
            last_name='Learner',
        )
        self.client.force_login(learner)
        profile_url = reverse('student_view_profile')
        response = self.client.post(profile_url, {
            'first_name': learner.first_name,
            'last_name': learner.last_name,
            'current_password': 'wrong-current-password',
            'new_password': 'A-Different-Strong-Password-789!',
            'confirm_password': 'A-Different-Strong-Password-789!',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Current password is incorrect')
        learner.refresh_from_db()
        self.assertTrue(learner.check_password('Original-Strong-Password-789!'))

        response = self.client.post(profile_url, {
            'first_name': learner.first_name,
            'last_name': learner.last_name,
            'current_password': 'Original-Strong-Password-789!',
            'new_password': 'A-Different-Strong-Password-789!',
            'confirm_password': 'A-Different-Strong-Password-789!',
        })
        self.assertRedirects(
            response, profile_url, fetch_redirect_response=False)
        learner.refresh_from_db()
        self.assertTrue(learner.check_password('A-Different-Strong-Password-789!'))
        self.assertEqual(self.client.get(profile_url).status_code, 200)

    def test_profile_picture_can_be_updated_for_both_roles(self):
        learner = CustomUser.objects.create_user(
            email='picture-learner@example.com',
            password='valid-password',
            user_type='2',
            gender='F',
            address='',
        )
        admin = CustomUser.objects.create_user(
            email='picture-admin@example.com',
            password='valid-password',
            user_type='1',
            gender='F',
            address='',
        )
        with tempfile.TemporaryDirectory() as media_root:
            with override_settings(MEDIA_ROOT=media_root):
                for user, profile_url in (
                        (learner, reverse('student_view_profile')),
                        (admin, reverse('admin_view_profile'))):
                    with self.subTest(role=user.user_type):
                        self.client.force_login(user)
                        response = self.client.post(profile_url, {
                            'first_name': 'Updated',
                            'last_name': 'Picture',
                            'profile_pic': self._profile_image_upload(),
                        })
                        self.assertRedirects(
                            response, profile_url, fetch_redirect_response=False)
                        user.refresh_from_db()
                        self.assertTrue(user.profile_pic.name)
                        self.assertTrue(user.profile_pic.name.endswith('.png'))

    @staticmethod
    def _profile_image_upload():
        image_data = BytesIO()
        Image.new('RGB', (2, 2), color='white').save(image_data, format='PNG')
        return SimpleUploadedFile(
            'profile.png', image_data.getvalue(), content_type='image/png')


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class PINAndLockoutTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = CustomUser.objects.create_user(
            username='pin-teacher',
            password='A-Strong-Password-283!',
            user_type='3',
            gender='F',
            first_name='PIN',
            last_name='Teacher',
        )
        self.legacy_learner = CustomUser.objects.create_user(
            username='legacy-learner-pin',
            password='A-Strong-Password-283!',
            user_type='2',
            gender='F',
        )
        issue_pin(self.legacy_learner)
        self.pin = issue_pin(self.user)

    def _sign_in(self, **overrides):
        data = {
            'username': self.user.username,
            'pin': self.pin,
            'role': 'teacher',
        }
        data.update(overrides)
        return self.client.post(reverse('user_login'), data)

    def test_staff_sign_in_uses_a_hashed_pin_without_a_password(self):
        self.user.refresh_from_db()
        self.assertNotEqual(self.user.pin_hash, self.pin)
        self.assertTrue(check_password(self.pin, self.user.pin_hash))
        self.assertFalse(self.user.has_usable_password())
        response = self._sign_in(pin='')
        self.assertRedirects(
            response,
            reverse('role_login', args=['teacher']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

        response = self._sign_in()
        self.assertRedirects(
            response,
            reverse('teacher_home'),
            fetch_redirect_response=False,
        )
        self.user.refresh_from_db()
        self.assertEqual(self.user.failed_login_attempts, 0)
        self.assertIsNone(self.user.locked_until)

    def test_pin_generator_rejects_weak_values(self):
        with patch(
            'main_app.credentials.secrets.randbelow',
            side_effect=[0, 11111, 12345, 54321, 22222, 14256],
        ):
            self.assertEqual(generate_pin(), '14256')

    def test_validate_pin_rejects_wrong_length_and_sequences(self):
        with self.assertRaises(ValueError):
            validate_pin('12345')
        with self.assertRaises(ValueError):
            validate_pin('1234')
        self.assertEqual(validate_pin('90817'), '90817')

    def test_all_sign_in_failures_use_the_same_generic_error(self):
        attempts = (
            ('missing-user', self.pin, 'teacher'),
            (self.user.username, '24680', 'teacher'),
            (self.user.username, self.pin, 'admin'),
        )
        for username, pin, role in attempts:
            with self.subTest(username=username, role=role, pin=pin):
                response = self.client.post(reverse('user_login'), {
                    'username': username,
                    'pin': pin,
                    'role': role,
                })
                self.assertRedirects(
                    response,
                    reverse('role_login', args=[role]),
                    fetch_redirect_response=False,
                )
                error_page = self.client.get(
                    reverse('role_login', args=[role]))
                self.assertContains(error_page, 'Invalid sign-in details')
                self.assertNotIn('_auth_user_id', self.client.session)

    @override_settings(
        SIMS_LOGIN_MAX_FAILED_ATTEMPTS=5,
        SIMS_LOGIN_LOCKOUT_MINUTES=15,
    )
    def test_five_failed_attempts_lock_the_account_and_success_unlocks_it(self):
        for _ in range(5):
            self._sign_in(pin='24680')

        self.user.refresh_from_db()
        self.assertEqual(self.user.failed_login_attempts, 5)
        self.assertGreater(self.user.locked_until, timezone.now())

        response = self._sign_in()
        self.assertRedirects(
            response,
            reverse('role_login', args=['teacher']),
            fetch_redirect_response=False,
        )
        self.assertNotIn('_auth_user_id', self.client.session)

        self.user.locked_until = timezone.now() - timedelta(seconds=1)
        self.user.save(update_fields=['locked_until'])
        response = self._sign_in()
        self.assertRedirects(
            response,
            reverse('teacher_home'),
            fetch_redirect_response=False,
        )
        self.user.refresh_from_db()
        self.assertEqual(self.user.failed_login_attempts, 0)
        self.assertIsNone(self.user.locked_until)

    @override_settings(SIMS_LOGIN_IP_MAX_FAILURES=2)
    def test_ip_failure_limit_blocks_a_later_correct_sign_in(self):
        self._sign_in(pin='24680')
        self._sign_in(pin='24680')
        response = self._sign_in()

        self.assertRedirects(
            response,
            reverse('role_login', args=['teacher']),
            fetch_redirect_response=False,
        )
        error_page = self.client.get(reverse('role_login', args=['teacher']))
        self.assertContains(error_page, 'Invalid sign-in details')
        self.assertNotIn('_auth_user_id', self.client.session)
        self.user.refresh_from_db()
        self.assertEqual(self.user.failed_login_attempts, 2)

    def test_staff_without_a_pin_cannot_sign_in(self):
        account = CustomUser.objects.create_user(
            username='unissued-pin',
            password=None,
            user_type='3',
            gender='F',
        )
        response = self.client.post(reverse('user_login'), {
            'username': account.username,
            'pin': '',
            'role': 'teacher',
        })
        self.assertRedirects(
            response,
            reverse('role_login', args=['teacher']),
            fetch_redirect_response=False,
        )
        error_page = self.client.get(reverse('role_login', args=['teacher']))
        self.assertContains(error_page, 'Invalid sign-in details')

    def test_issue_pins_command_creates_hashes_and_one_time_csv(self):
        second = CustomUser.objects.create_user(
            username='another-user',
            password='A-Strong-Password-283!',
            user_type='3',
            gender='M',
            first_name='Another',
            last_name='Teacher',
        )

        with tempfile.TemporaryDirectory() as directory:
            csv_path = f'{directory}\\pins.csv'
            output = StringIO()
            call_command('issue_pins', csv_path, stdout=output)

            with open(csv_path, newline='', encoding='utf-8') as csv_file:
                rows = list(csv.reader(csv_file))
            self.assertEqual(rows[0], ['name', 'role', 'username', 'PIN'])
            self.assertEqual(len(rows), 3)
            self.assertIn('delete the CSV', output.getvalue())
            self.assertNotIn(self.pin, output.getvalue())

        second.refresh_from_db()
        self.assertTrue(second.pin_hash)
        self.assertFalse(second.has_usable_password())
        self.assertTrue(check_password(rows[2][3], second.pin_hash))
        self.legacy_learner.refresh_from_db()
        self.assertFalse(self.legacy_learner.pin_hash)


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class CredentialResetTests(TestCase):
    def setUp(self):
        self.admin = CustomUser.objects.create_user(
            username='reset-admin',
            password='Original-Admin-Password-381!',
            user_type='1',
            gender='M',
        )
        self.teacher = CustomUser.objects.create_user(
            username='reset-teacher',
            password='Original-Teacher-Password-381!',
            user_type='3',
            gender='F',
            first_name='Taylor',
            last_name='Teacher',
        )
        self.learner = CustomUser.objects.create_user(
            username='reset-learner',
            password='Original-Learner-Password-381!',
            user_type='2',
            gender='F',
            first_name='Casey',
            last_name='Learner',
        )
        self.learner.student.teacher = self.teacher.teacher
        self.learner.student.save()
        self.admin_pin = issue_pin(self.admin)
        self.teacher_pin = issue_pin(self.teacher)

    def _reset_url(self, user):
        return reverse('reset_user_credentials', args=[user.pk])

    def test_administrator_can_reset_learner_password(self):
        self.client.force_login(self.admin)
        new_password = 'Different-Strong-Password-932!'
        self.learner.failed_login_attempts = 4
        self.learner.locked_until = timezone.now() + timedelta(minutes=15)
        response = self.client.post(self._reset_url(self.learner), {
            'action': 'password',
            'new_password': new_password,
            'confirm_password': new_password,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, new_password)
        self.learner.refresh_from_db()
        self.assertTrue(self.learner.check_password(new_password))
        self.assertEqual(self.learner.failed_login_attempts, 0)
        self.assertIsNone(self.learner.locked_until)
        self.assertFalse(self.teacher.has_usable_password())

    def test_administrator_and_teacher_can_reset_allowed_pins(self):
        self.client.force_login(self.admin)
        response = self.client.post(self._reset_url(self.teacher), {
            'action': 'pin',
        })
        self.assertEqual(response.status_code, 200)
        new_teacher_pin = response.context['secret']
        self.teacher.refresh_from_db()
        self.assertNotEqual(new_teacher_pin, self.teacher_pin)
        self.assertTrue(check_password(new_teacher_pin, self.teacher.pin_hash))
        self.assertFalse(self.teacher.has_usable_password())

        self.client.force_login(self.teacher)
        response = self.client.post(self._reset_url(self.learner), {
            'action': 'password',
            'new_password': 'Different-Learner-Password-932!',
            'confirm_password': 'Different-Learner-Password-932!',
        })
        self.assertEqual(response.status_code, 200)
        self.learner.refresh_from_db()
        self.assertTrue(self.learner.check_password(
            'Different-Learner-Password-932!'))
        self.assertFalse(self.learner.pin_hash)

    def test_teacher_cannot_reset_teacher_or_administrator(self):
        self.client.force_login(self.teacher)
        response = self.client.get(self._reset_url(self.teacher))
        self.assertEqual(response.status_code, 403)
        response = self.client.get(self._reset_url(self.admin))
        self.assertEqual(response.status_code, 404)
        unassigned_learner = CustomUser.objects.create_user(
            username='unassigned-reset-target',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        response = self.client.get(self._reset_url(unassigned_learner))
        self.assertEqual(response.status_code, 404)

    def test_learner_can_change_password_without_a_pin(self):
        self.client.force_login(self.learner)
        response = self.client.post(reverse('student_view_profile'), {
            'first_name': self.learner.first_name,
            'last_name': self.learner.last_name,
            'current_password': 'Original-Learner-Password-381!',
            'new_password': 'Changed-Learner-Password-853!',
            'confirm_password': 'Changed-Learner-Password-853!',
        })
        self.assertRedirects(
            response, reverse('student_view_profile'),
            fetch_redirect_response=False)
        self.learner.refresh_from_db()
        self.assertTrue(
            self.learner.check_password('Changed-Learner-Password-853!'))
        self.assertEqual(self.client.get(
            reverse('student_view_profile')).status_code, 200)

    def test_teacher_can_reset_own_pin_with_the_current_pin(self):
        self.client.force_login(self.teacher)
        response = self.client.post(reverse('teacher_profile'), {
            'first_name': self.teacher.first_name,
            'last_name': self.teacher.last_name,
            'current_pin': self.teacher_pin,
            'reset_pin': 'on',
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn('issued_pin', response.context)
        self.teacher.refresh_from_db()
        self.assertFalse(self.teacher.has_usable_password())

    def test_admin_password_and_pin_have_no_web_reset_page(self):
        self.client.force_login(self.admin)
        response = self.client.get(self._reset_url(self.admin))
        self.assertEqual(response.status_code, 404)
        self.client.logout()
        self.admin.is_staff = True
        self.admin.is_superuser = True
        self.admin.save(update_fields=['is_staff', 'is_superuser'])
        self.client.force_login(self.admin)
        for path in (
            '/admin/password_change/',
            '/admin/password_change/done/',
            f'/admin/main_app/customuser/{self.admin.pk}/password/',
        ):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 404)

    def test_email_password_reset_and_email_availability_urls_are_removed(self):
        for path in (
            '/accounts/password_reset/',
            '/accounts/password_reset/done/',
            '/accounts/reset/uid/token/',
            '/accounts/reset/done/',
            '/check_email_availability',
        ):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 404)
        response = self.client.get(reverse('role_login', args=['learner']))
        self.assertNotContains(response, 'password_reset')
        self.assertContains(
            response,
            'Forgot your password? Ask your teacher or administrator to reset it.')

    def test_operator_commands_reset_admin_credentials_without_crossing_them(self):
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            call_command(
                'set_admin_password',
                self.admin.username,
                stdout=StringIO(),
            )

        django_superuser = CustomUser.objects.create_superuser(
            username='django-superuser',
            password='Original-Superuser-Password-923!',
            user_type='1',
            gender='M',
        )
        with patch(
            'getpass.getpass',
            side_effect=[
                'New-Operator-Password-625!',
                'New-Operator-Password-625!',
            ],
        ):
            output = StringIO()
            call_command(
                'set_admin_password',
                django_superuser.username,
                stdout=output,
            )
        django_superuser.refresh_from_db()
        self.assertTrue(django_superuser.check_password(
            'New-Operator-Password-625!'))
        self.assertNotIn('New-Operator-Password-625!', output.getvalue())

        output = StringIO()
        call_command(
            'reset_admin_pin',
            self.admin.username,
            stdout=output,
        )
        admin_pin = output.getvalue().strip().split(': ')[-1]
        self.admin.refresh_from_db()
        self.assertTrue(check_password(admin_pin, self.admin.pin_hash))
        self.assertFalse(self.admin.has_usable_password())


class ActivityLogTests(TestCase):
    def setUp(self):
        self.admin = CustomUser.objects.create_user(
            username='activity-admin',
            password='Valid-Admin-Password-389!',
            user_type='1',
            gender='F',
        )

    def test_model_changes_include_actor_and_assignment_values_without_secrets(self):
        request = RequestFactory().post('/admin/assignments/')
        request.user = self.admin
        token = set_current_request(request)
        try:
            teacher_user = CustomUser.objects.create_user(
                username='activity-teacher',
                password='Never-Log-This-Password-391!',
                user_type='3',
                gender='M',
                first_name='Casey',
                last_name='Teacher',
            )
            teacher_pin = issue_pin(teacher_user)
            learner_user = CustomUser.objects.create_user(
                username='activity-learner',
                password='Never-Log-Learner-Password-392!',
                user_type='2',
                gender='F',
            )
            learner = learner_user.student
            learner.teacher = teacher_user.teacher
            learner.save(update_fields=['teacher', 'updated_at'])
        finally:
            reset_current_request(token)

        assignment = ActivityLog.objects.filter(
            object_type='main_app.student',
            object_id=str(learner.pk),
            action='update',
        ).latest('created_at')
        self.assertEqual(assignment.actor_id, self.admin.pk)
        self.assertEqual(assignment.target_user_id, learner_user.pk)
        self.assertIsNone(assignment.changes['teacher']['old'])
        self.assertIn(
            'Casey Teacher', assignment.changes['teacher']['new'])

        user_events = ActivityLog.objects.filter(
            object_type='main_app.customuser',
            object_id=str(teacher_user.pk),
        )
        serialized = ' '.join(str(event.changes) for event in user_events)
        self.assertNotIn('Never-Log-This-Password-391!', serialized)
        self.assertNotIn(teacher_pin, serialized)
        self.assertIn('[redacted]', serialized)

    def test_admin_can_filter_export_and_view_a_user_timeline(self):
        learner = CustomUser.objects.create_user(
            username='activity-target',
            password='Valid-Learner-Password-391!',
            user_type='2',
            gender='F',
        )
        event = ActivityLog.objects.create(
            actor=self.admin,
            actor_username=self.admin.username,
            target_user=learner,
            action='update',
            object_type='main_app.student',
            object_id=str(learner.student.pk),
            description='Updated activity target learner',
            changes={'teacher': {'old': None, 'new': 'Assigned Teacher'}},
        )
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('activity_log'),
            {'search': 'activity target', 'action': 'update'},
        )
        self.assertContains(response, event.description)

        response = self.client.get(
            reverse('activity_log'), {'action': 'update', 'export': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertContains(response, 'Updated activity target learner')

        response = self.client.get(
            reverse('activity_user_timeline', args=[learner.pk]))
        self.assertContains(response, event.description)

        teacher = CustomUser.objects.create_user(
            username='activity-nonadmin',
            password='Valid-Teacher-Password-391!',
            user_type='3',
            gender='M',
        )
        self.client.force_login(teacher)
        response = self.client.get(reverse('activity_log'))
        self.assertRedirects(
            response, reverse('teacher_home'),
            fetch_redirect_response=False)

    def test_deleting_target_account_keeps_a_valid_user_snapshot(self):
        learner = CustomUser.objects.create_user(
            username='activity-deleted-target',
            password='Valid-Learner-Password-392!',
            user_type='2',
            gender='F',
        )
        learner_id = str(learner.pk)
        learner.delete()

        deletion_events = ActivityLog.objects.filter(
            target_user_id_snapshot=learner_id,
            action='delete',
        )
        self.assertGreaterEqual(deletion_events.count(), 2)
        self.assertTrue(all(
            event.target_user_id is None
            for event in deletion_events
        ))
        self.assertTrue(all(
            event.target_username == 'activity-deleted-target'
            for event in deletion_events
        ))

    @override_settings(SIMS_ACTIVITY_LOG_RETENTION_DAYS=30)
    def test_retention_command_removes_only_entries_past_the_cutoff(self):
        old_entry = ActivityLog.objects.create(
            actor_username='old-user',
            action='delete',
            object_type='main_app.student',
            description='Old entry',
        )
        recent_entry = ActivityLog.objects.create(
            actor_username='recent-user',
            action='create',
            object_type='main_app.student',
            description='Recent entry',
        )
        ActivityLog.objects.filter(pk=old_entry.pk).update(
            created_at=timezone.now() - timedelta(days=31))

        call_command('prune_activity_log', stdout=StringIO())

        self.assertFalse(ActivityLog.objects.filter(pk=old_entry.pk).exists())
        self.assertTrue(ActivityLog.objects.filter(pk=recent_entry.pk).exists())


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class DefaultAccountCommandTests(TestCase):
    @staticmethod
    def _create_accounts(**options):
        output = StringIO()
        call_command('create_default_accounts', stdout=output, **options)
        return output.getvalue()

    def test_command_creates_distinct_role_accounts_with_generated_secrets(self):
        output = self._create_accounts()
        accounts = {
            account.username: account
            for account in CustomUser.objects.filter(
                username__in=(
                    'mbangulap-admin',
                    'mbangulap-teacher',
                    'mbangulap',
                )
            )
        }
        self.assertEqual(len(accounts), 3)
        admin = accounts['mbangulap-admin']
        teacher = accounts['mbangulap-teacher']
        learner = accounts['mbangulap']
        self.assertEqual(admin.user_type, '1')
        self.assertEqual(teacher.user_type, '3')
        self.assertEqual(learner.user_type, '2')
        self.assertFalse(admin.has_usable_password())
        self.assertFalse(teacher.has_usable_password())
        self.assertTrue(learner.has_usable_password())
        self.assertTrue(learner.check_password(
            output.split('Learner: username=mbangulap, password=', 1)[1]
            .splitlines()[0]))
        self.assertIn('Administrator: username=mbangulap-admin, PIN=', output)
        self.assertIn('Teacher: username=mbangulap-teacher, PIN=', output)
        admin_pin = output.split(
            'Administrator: username=mbangulap-admin, PIN=', 1)[1].splitlines()[0]
        teacher_pin = output.split(
            'Teacher: username=mbangulap-teacher, PIN=', 1)[1].splitlines()[0]
        self.assertNotEqual(admin_pin, teacher_pin)
        self.assertEqual(
            [admin.is_superuser, teacher.is_superuser, learner.is_superuser],
            [False, False, False],
        )

    def test_command_refuses_to_overwrite_existing_account(self):
        CustomUser.objects.create_user(
            username='mbangulap',
            password='A-Valid-Learner-Password-927!',
            user_type='2',
            gender='F',
        )
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            call_command('create_default_accounts', stdout=StringIO())
        self.assertFalse(
            CustomUser.objects.filter(username='mbangulap-admin').exists())

    def test_command_can_reuse_hidden_credentials_across_environments(self):
        with patch(
            'main_app.management.commands.create_default_accounts.getpass.getpass',
            side_effect=[
                'Strong-Shared-Learner-Pass-782!',
                '48291',
                '73914',
            ],
        ):
            output = self._create_accounts(reuse_credentials=True)

        self.assertIn(
            'Administrator: username=mbangulap-admin, PIN=48291', output)
        self.assertIn(
            'Teacher: username=mbangulap-teacher, PIN=73914', output)
        learner = CustomUser.objects.get(username='mbangulap')
        self.assertTrue(learner.check_password(
            'Strong-Shared-Learner-Pass-782!'))


@override_settings(PASSWORD_HASHERS=[
    'django.contrib.auth.hashers.MD5PasswordHasher',
])
class LearnerImportAndExportTests(TestCase):
    def setUp(self):
        self.admin = CustomUser.objects.create_user(
            username='export-admin',
            password='valid-password',
            user_type='1',
            gender='F',
        )
        self.teacher_user = CustomUser.objects.create_user(
            username='export-teacher',
            password='valid-password',
            user_type='3',
            gender='M',
        )
        self.period = AcademicPeriod.objects.create(
            year=2026, term=1, is_current=True)
        self.course = Course.objects.create(name='Grade 7')
        self.client.force_login(self.admin)

    @staticmethod
    def _csv_file(rows, name='learners.csv'):
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow([
            'first_name', 'last_name', 'username', 'gender', 'course',
        ])
        writer.writerows(rows)
        return SimpleUploadedFile(
            name, buffer.getvalue().encode('utf-8'), content_type='text/csv')

    def _preview_csv(self, rows, name='learners.csv'):
        return self.client.post(
            reverse('import_learners'),
            {
                'action': 'upload',
                'file': self._csv_file(rows, name=name),
            },
        )

    def test_csv_import_previews_then_creates_learners_with_one_time_passwords(self):
        response = self._preview_csv([[
            'Ada', 'Okafor', 'ada-imported', 'F', 'Grade 7',
        ]])

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Preview (1 learner rows)')
        self.assertContains(response, 'Ready')
        token = response.context['preview_token']

        response = self.client.post(reverse('import_learners'), {
            'action': 'confirm',
            'preview_token': token,
            'academic_period': self.period.pk,
            'teacher': self.teacher_user.teacher.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn('no-store', response['Cache-Control'])
        self.assertContains(response, 'shown only once')
        learner = CustomUser.objects.get(username='ada-imported')
        self.assertEqual(learner.user_type, '2')
        self.assertEqual(learner.student.course, self.course)
        self.assertEqual(learner.student.academic_period, self.period)
        self.assertEqual(learner.student.teacher, self.teacher_user.teacher)
        self.assertIsNone(learner.student.created_by)
        password = response.context['credentials'][0]['password']
        self.assertTrue(learner.check_password(password))
        self.assertTrue(ActivityLog.objects.filter(action='import').exists())
        self.assertNotIn('learner_import_preview', self.client.session)

    def test_teacher_import_assigns_accounts_only_to_themselves(self):
        self.client.force_login(self.teacher_user)
        response = self._preview_csv([[
            'Kai', 'Learner', 'kai-imported', 'M', 'Grade 7',
        ]])
        token = response.context['preview_token']
        response = self.client.post(reverse('import_learners'), {
            'action': 'confirm',
            'preview_token': token,
            'academic_period': self.period.pk,
        })

        learner = CustomUser.objects.get(username='kai-imported')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(learner.student.teacher, self.teacher_user.teacher)
        self.assertEqual(learner.student.created_by, self.teacher_user.teacher)

    def test_invalid_duplicate_and_unknown_course_rows_cannot_be_imported(self):
        response = self._preview_csv([
            ['Ada', 'Okafor', 'duplicate-import', 'F', 'Grade 7'],
            ['Bea', 'Bello', 'duplicate-import', 'F', 'Missing grade'],
        ])

        self.assertTrue(response.context['preview_has_errors'])
        self.assertContains(response, 'duplicated in this file')
        self.assertContains(response, 'does not exist')
        response = self.client.post(reverse('import_learners'), {
            'action': 'confirm',
            'preview_token': response.context['preview_token'],
            'academic_period': self.period.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(CustomUser.objects.filter(
            username='duplicate-import').exists())
        self.assertContains(response, 'No accounts will be created')

    @override_settings(SIMS_LEARNER_IMPORT_MAX_ROWS=1)
    def test_import_row_limit_is_enforced(self):
        response = self._preview_csv([
            ['Ada', 'Okafor', 'first-import', 'F', 'Grade 7'],
            ['Bea', 'Bello', 'second-import', 'F', 'Grade 7'],
        ])
        self.assertContains(response, 'exceeds the 1-row limit')
        self.assertFalse(response.context['preview_rows'])

    def test_xlsx_import_and_template_are_supported(self):
        from openpyxl import Workbook, load_workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            'first_name', 'last_name', 'username', 'gender', 'course',
        ])
        worksheet.append(['Sam', 'Learner', 'sam-xlsx', 'M', 'Grade 7'])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'learners.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'),
        )
        response = self.client.post(reverse('import_learners'), {
            'action': 'upload',
            'file': upload,
        })
        self.assertContains(response, 'sam-xlsx')

        template = self.client.get(
            reverse('download_learner_import_template'),
            {'format': 'xlsx'},
        )
        self.assertEqual(template.status_code, 200)
        template_book = load_workbook(BytesIO(template.content), read_only=True)
        self.assertEqual(
            list(next(template_book.active.iter_rows(values_only=True))),
            ['first_name', 'last_name', 'username', 'gender', 'course'],
        )
        template_book.close()

    def test_import_preview_expires_and_learner_role_is_denied(self):
        response = self._preview_csv([[
            'Ada', 'Okafor', 'expired-import', 'F', 'Grade 7',
        ]])
        token = response.context['preview_token']
        session = self.client.session
        preview = session['learner_import_preview']
        preview['created_at'] = timezone.now().timestamp() - 901
        session['learner_import_preview'] = preview
        session.save()
        response = self.client.post(reverse('import_learners'), {
            'action': 'confirm',
            'preview_token': token,
            'academic_period': self.period.pk,
        })
        self.assertRedirects(
            response, reverse('import_learners'),
            fetch_redirect_response=False)
        self.assertFalse(CustomUser.objects.filter(
            username='expired-import').exists())

        learner = CustomUser.objects.create_user(
            username='import-denied',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        self.client.force_login(learner)
        response = self.client.get(reverse('import_learners'))
        self.assertRedirects(
            response, reverse('student_home'),
            fetch_redirect_response=False)

    def test_learner_csv_export_sanitizes_formulas_and_logs_the_export(self):
        learner = CustomUser.objects.create_user(
            username='csv-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            first_name='=1+1',
        )
        learner.student.course = self.course
        learner.student.academic_period = self.period
        learner.student.save()

        response = self.client.get(reverse('manage_student'), {'export': 'csv'})

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn(b"'=1+1", response.content)
        self.assertIn('learner-list-term-1-2026.csv', response['Content-Disposition'])
        self.assertTrue(ActivityLog.objects.filter(
            action='export', actor=self.admin,
            description__contains='Learner List').exists())

    def test_teacher_list_and_login_activity_support_pdf_and_csv(self):
        teacher_csv = self.client.get(
            reverse('manage_teachers'), {'export': 'csv'})
        self.assertEqual(teacher_csv['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn(b'export-teacher', teacher_csv.content)
        teacher_pdf = self.client.get(
            reverse('manage_teachers'), {'export': 'pdf'})
        self.assertTrue(teacher_pdf.content.startswith(b'%PDF'))

        learner = CustomUser.objects.create_user(
            username='login-export-learner',
            password='valid-password',
            user_type='2',
            gender='M',
        )
        learner.student.course = self.course
        learner.student.save()
        LoginActivity.objects.create(student=learner.student)
        login_csv = self.client.get(
            reverse('learner_login_activity'), {'export': 'csv'})
        self.assertEqual(login_csv.status_code, 200)
        self.assertIn(b'login-export-learner', login_csv.content)
        login_pdf = self.client.get(
            reverse('learner_login_activity'), {'export': 'pdf'})
        self.assertTrue(login_pdf.content.startswith(b'%PDF'))

        learner_pdf = self.client.get(
            reverse('manage_student'), {'export': 'pdf'})
        self.assertTrue(learner_pdf.content.startswith(b'%PDF'))

    def test_activity_log_report_and_report_cards_export_as_csv(self):
        ActivityLog.objects.create(
            actor=self.admin,
            actor_username=self.admin.username,
            action='sign_in',
            description='Export test activity',
        )
        activity_csv = self.client.get(
            reverse('activity_log'), {'export': 'csv'})
        self.assertEqual(activity_csv.status_code, 200)
        self.assertIn(b'Export test activity', activity_csv.content)
        activity_pdf = self.client.get(
            reverse('activity_log'), {'export': 'pdf'})
        self.assertTrue(activity_pdf.content.startswith(b'%PDF'))

        learner = CustomUser.objects.create_user(
            username='report-export-learner',
            password='valid-password',
            user_type='2',
            gender='F',
            first_name='Report',
            last_name='Learner',
        )
        learner.student.course = self.course
        learner.student.academic_period = self.period
        learner.student.teacher = self.teacher_user.teacher
        learner.student.save()
        subject = Subject.objects.create(
            name='Mathematics',
            course=self.course,
            teacher=self.teacher_user.teacher,
        )
        StudentResult.objects.create(
            student=learner.student, subject=subject, test=20, exam=70)
        report_csv = self.client.post(reverse('generate_report'), {
            'report_type': 'school_summary',
            'export_format': 'csv',
        })
        self.assertEqual(report_csv.status_code, 200)
        self.assertIn(b'Total learners', report_csv.content)

        self.client.force_login(self.teacher_user)
        report_card_csv = self.client.post(
            reverse('generate_report_card'), {
                'student': learner.student.pk,
                'academic_period': self.period.pk,
                'export_format': 'csv',
            })
        self.assertEqual(report_card_csv.status_code, 200)
        self.assertIn(b'Mathematics', report_card_csv.content)

        self.client.force_login(learner)
        own_card_csv = self.client.get(
            reverse('student_report_card'), {'export': 'csv'})
        self.assertEqual(own_card_csv.status_code, 200)
        self.assertIn(b'Mathematics', own_card_csv.content)

    def test_empty_exports_return_page_message_without_a_file(self):
        response = self.client.get(
            reverse('learner_login_activity'), {'export': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/html; charset=utf-8')
        self.assertContains(response, 'No login activity matches these filters.')

    def test_staff_forms_default_to_the_current_academic_period(self):
        from .forms import (
            AttendanceSelectionForm,
            ReportCardForm,
            ReportGenerationForm,
            StudentForm,
        )

        self.assertEqual(
            StudentForm().fields['academic_period'].initial, self.period)
        self.assertEqual(
            ReportGenerationForm().fields['academic_period'].initial,
            self.period,
        )
        self.assertEqual(
            ReportCardForm(user=self.teacher_user).fields[
                'academic_period'].initial,
            self.period,
        )
        self.assertEqual(
            AttendanceSelectionForm().fields['academic_period'].initial,
            self.period,
        )

    def test_admin_can_change_current_period_and_used_period_cannot_be_deleted(self):
        next_period = AcademicPeriod.objects.create(year=2026, term=2)
        response = self.client.post(reverse(
            'set_current_academic_period', args=[next_period.pk]))
        self.assertRedirects(
            response, reverse('manage_academic_periods'),
            fetch_redirect_response=False)
        self.period.refresh_from_db()
        next_period.refresh_from_db()
        self.assertFalse(self.period.is_current)
        self.assertTrue(next_period.is_current)

        learner = CustomUser.objects.create_user(
            username='period-delete-learner',
            password='valid-password',
            user_type='2',
            gender='F',
        )
        learner.student.academic_period = next_period
        learner.student.save()
        subject = Subject.objects.create(name='Period subject')
        StudentResult.objects.create(
            student=learner.student,
            subject=subject,
            test=10,
            exam=20,
        )
        response = self.client.post(reverse(
            'delete_academic_period', args=[next_period.pk]))
        self.assertRedirects(
            response, reverse('manage_academic_periods'),
            fetch_redirect_response=False)
        self.assertTrue(AcademicPeriod.objects.filter(pk=next_period.pk).exists())
        self.assertContains(
            self.client.get(reverse('manage_academic_periods')),
            'cannot be deleted',
        )

    def test_xlsx_formula_cells_and_oversized_uploads_are_rejected(self):
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            'first_name', 'last_name', 'username', 'gender', 'course',
        ])
        worksheet.append(['=1+1', 'Learner', 'formula-xlsx', 'M', 'Grade 7'])
        buffer = BytesIO()
        workbook.save(buffer)
        formula_file = SimpleUploadedFile(
            'formula.xlsx', buffer.getvalue())
        response = self.client.post(reverse('import_learners'), {
            'action': 'upload',
            'file': formula_file,
        })
        self.assertContains(response, 'Formula cells are not allowed')
        self.assertFalse(response.context['preview_rows'])

        with override_settings(SIMS_LEARNER_IMPORT_MAX_BYTES=4):
            oversized = SimpleUploadedFile(
                'too-large.csv',
                b'first_name,last_name,username,gender,course',
            )
            response = self.client.post(reverse('import_learners'), {
                'action': 'upload',
                'file': oversized,
            })
        self.assertContains(response, 'exceeds the')
